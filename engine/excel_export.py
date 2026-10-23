"""Export the current report results as formatted Excel tables."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from .rules import display_order_key
from .render_backup import floor_calculation, missing_metric_reason


def export_report(path, kind, data):
    book = Workbook()
    book.remove(book.active)
    font = Font(name='Calibri', size=11)
    alignment = Alignment(vertical='top', wrap_text=True)
    header_fill = PatternFill('solid', fgColor='245E6B')
    header_font = Font(name='Calibri', bold=True, color='FFFFFF')

    def sheet(name, rows):
        ws = book.create_sheet(name)
        for row in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub('', x) if isinstance(x, str) else x for x in row])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = font
                cell.alignment = alignment
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.##;[Red]-0.##'
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 30 if i < 4 else 48
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    if kind == 'ladder':
        sheet('配置阶梯', [['车型', '版型', '指导价（万元）', '比较基准', '基础配置或差异配置', '选装配置']] + [
            [data['model'], c['name'], c.get('price'), c.get('base') or '基础配置',
             '\n'.join(c['items']), '\n'.join(c['options'])] for c in data['columns']])
    elif kind == 'diff':
        groups = data['groups']
        left, right = data.get('self_model') or '本品', data.get('comp_model') or '竞品'
        def header(g):
            def price(p): return '待核' if p is None else f'{p:g}万'
            return f"{g['pair']['self_trim']} {price(g.get('self_price'))}\nVS\n{g['pair']['comp_trim']} {price(g.get('comp_price'))}"
        rows = [[f'{left} vs {right} · 竞争力对比'], ['版型配对'] + [header(g) for g in groups],
                [left+'多'] + ['\n'.join(g['more']) or '—' for g in groups],
                [left+'少'] + ['\n'.join(g['less']) or '—' for g in groups]]
        for key, label in [('config_adv','配置优势（元）'), ('flat_adv','拉平指导价优势（元）'), ('overall','综合竞争力（元）')]:
            rows.append([label] + [g.get('valuation', {}).get(key) if g.get('valuation', {}).get(key) is not None else
                                  missing_metric_reason(g, key) for g in groups])
        if any(g.get('chip_comparison') for g in groups):
            rows.append(['车机芯片（暂不赋值）'] + [g.get('chip_comparison') or '—' for g in groups])
        for field, model in (('self_options', left), ('comp_options', right)):
            if any(g.get(field) for g in groups):
                rows.append([model+'选装（仅展示，不参与赋值）'] + [
                    '\n'.join(g.get(field, []) or []) or '—' for g in groups])
        sheet('对比汇总', rows)
        ws = book['对比汇总']
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(groups)+1)
        ws.auto_filter.ref = None
        ws.freeze_panes = 'B3'
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 26
        for col in range(2, len(groups)+2):
            ws.column_dimensions[get_column_letter(col)].width = 46
        border = Border(*([Side(style='thin', color='DDE7EB')]*4))
        colors = {1: ('EDF4F5', '245E6B'), 3: ('EFF8F1', '28744B'),
                  4: ('FFF2F1', 'B64D43'), 0: ('FFFFFF', '243444')}
        fills = {key: PatternFill('solid', fgColor=bg) for key, (bg, _) in colors.items()}
        fonts = {key: Font(name='Calibri', size=11, bold=key == 1, color=fg)
                 for key, (_, fg) in colors.items()}
        alignments = {(header, first): Alignment(vertical='center' if header or first else 'top',
                                                horizontal='center' if header else 'left', wrap_text=True)
                      for header in (False, True) for first in (False, True)}
        for row in ws:
            for cell in row:
                cell.border = border
                if cell.row > 1:
                    style = 1 if cell.column == 1 or cell.row == 2 else cell.row if cell.row in (3, 4) else 0
                    cell.fill = fills[style]
                    cell.font = fonts[style]
                cell.alignment = alignments[cell.row <= 2, cell.column == 1]
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 72
        for index in (3, 4, *range(8, ws.max_row + 1)):
            lines = max((sum(max(1,(len(line)+24)//25) for line in str(c.value).split('\n')) for c in ws[index]), default=1)
            ws.row_dimensions[index].height = min(409, max(60, lines*17+16))
        for index in (5,6,7): ws.row_dimensions[index].height = 30
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.print_area = ws.dimensions
        sheet('价格计算', [['左侧版型', '右侧版型', '左侧指导价（万元）', '右侧指导价（万元）',
                          '左侧底价（元）', '右侧底价（元）', '配置优势（元）', '拉平指导价优势（元）',
                          '综合竞争力（元）', '综合竞争力计算依据']] + [
            [g['pair']['self_trim'], g['pair']['comp_trim'], g.get('self_price'), g.get('comp_price'),
             g.get('self_floor_price'), g.get('comp_floor_price'), g.get('valuation', {}).get('config_adv'),
             g.get('valuation', {}).get('flat_adv'), g.get('valuation', {}).get('overall'),
             floor_calculation(g)]
            for g in groups])
        book['价格计算'].freeze_panes = 'C2'
        sheet('赋值明细', [['左侧版型', '右侧版型', '多或少', '配置差异', '计价依据', '金额（元）']] + [
            [g['pair']['self_trim'], g['pair']['comp_trim'], d['side'], d.get('display'), d.get('rule'), d.get('amount')]
            for g in groups for d in sorted(g.get('valuation', {}).get('detail', []), key=lambda item: display_order_key(item['no']))])
        ordered_nos = sorted({c['no'] for c in data['cells']}, key=display_order_key)
        serial_by_no = {no: index for index, no in enumerate(ordered_nos, 1)}
        sheet('配置判定', [['左侧版型', '右侧版型', '序号', '配置名称', '左侧配置', '右侧配置', '判定', '差异说明']] + [
            [groups[c['pair']]['pair']['self_trim'], groups[c['pair']]['pair']['comp_trim'], serial_by_no[c['no']],
             c.get('name'), c.get('self_val'), c.get('comp_val'), c.get('verdict'), c.get('display')]
            for c in sorted(data['cells'], key=lambda item: (item['pair'], display_order_key(item['no'])))])
    else:
        raise ValueError('不支持的导出类型')
    book.save(path)
