"""Seat summary line breaks are shared presentation, not new pricing items."""
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import differ, render_backup
from test_comparison_display import PAIR, ladder, result


SEAT_LINES = ['前排座椅通风加热按摩记忆', '主驾座椅头枕音响', '二排座椅通风加热']
FUNCTIONS = {
    **{seat + feature: '●' for seat in ('主驾', '副驾') for feature in ('通风', '加热', '按摩')},
    '主驾头枕音响': '●', '二排通风': '●', '二排加热': '●',
}
WITH_SEATS = {36: FUNCTIONS, 45: '前排座椅记忆',
              35: '主驾6向电调+副驾6向电调(主腰撑4向/副腰撑4向)'}
WITHOUT_SEATS = {35: '主驾6向电调+副驾4向电调'}


def test_summary_splits_seat_groups_but_preserves_other_configuration_text():
    directions = '副驾4向电调（副驾6向电调），无（(主腰撑4向/副腰撑4向)）'
    combined = '前排座椅通风加热按摩、主驾座椅头枕音响、二排座椅通风加热'
    assert differ._merge_seat_labels([
        '800V', combined, '前排座椅记忆', directions,
    ]) == ['800V', *SEAT_LINES, directions]
    # Duplicate memory does not create an extra line, and non-seat punctuation stays intact.
    assert differ._merge_seat_labels([
        *SEAT_LINES, '前排座椅记忆', '外后视镜电调、折叠', directions,
    ]) == [*SEAT_LINES, '外后视镜电调、折叠', directions]


@pytest.mark.parametrize('side', ['more', 'less'])
def test_real_comparison_splits_summary_without_changing_cells_or_valuation(side):
    left, right = (WITH_SEATS, WITHOUT_SEATS) if side == 'more' else (WITHOUT_SEATS, WITH_SEATS)
    cells, groups = result(left, right)
    original_cells, original_valuation = deepcopy(cells), deepcopy(groups[0]['valuation'])
    assert [text for text in groups[0][side] if text in SEAT_LINES] == SEAT_LINES
    assert len(groups[0]['valuation']['detail']) == 3  # Direction, functions and memory, not five lines.
    assert len(cells) == len(original_cells)
    render_backup.format_configurations(cells, groups)
    assert [text for text in groups[0][side] if text in SEAT_LINES] == SEAT_LINES
    direction = next(cell for cell in cells if cell['no'] == 35)
    assert direction['backup_' + side] in groups[0][side]
    assert '腰撑4向' not in direction['backup_' + side]
    for before, after in zip(original_cells, cells):
        if before['no'] in (36, 45):
            assert after == before
        else:
            assert {key: value for key, value in before.items() if key not in (
                'self_val', 'comp_val', 'display', 'backup_more', 'backup_less')} == {
                key: value for key, value in after.items() if key not in (
                    'self_val', 'comp_val', 'display', 'backup_more', 'backup_less')}
    assert {key: value for key, value in groups[0]['valuation'].items() if key != 'detail'} == {
        key: value for key, value in original_valuation.items() if key != 'detail'}
    assert [(item['no'], item['amount'], item['side']) for item in groups[0]['valuation']['detail']] == [
        (item['no'], item['amount'], item['side']) for item in original_valuation['detail']]


@pytest.mark.parametrize('side', ['more', 'less'])
def test_api_markdown_and_excel_use_identical_seat_line_breaks(tmp_path, side):
    left_values, right_values = (WITH_SEATS, WITHOUT_SEATS) if side == 'more' else (WITHOUT_SEATS, WITH_SEATS)
    paths = [tmp_path / 'left.json', tmp_path / 'right.json']
    ladder(left_values, '左版', '本品').save(paths[0])
    ladder(right_values, '右版', '竞品').save(paths[1])
    original_files = [path.read_bytes() for path in paths]
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        report = bridge.run_diff(*(str(path) for path in paths), [PAIR], left_is_competitor=True)
        assert report['ok'], report.get('error')
        group = report['groups'][0]
        assert [text for text in group[side] if text in SEAT_LINES] == SEAT_LINES
        label = '本品多' if side == 'more' else '本品少'
        assert '| ' + label + ' | ' + '<br>'.join(group[side]) + ' |' in report['md']
        assert Path(report['md_path']).read_text(encoding='utf-8') == report['md']
        exported = bridge.export_excel('diff', report)
        assert exported['ok'], exported.get('error')
        book = load_workbook(exported['path'])
        try:
            assert book['对比汇总']['B3' if side == 'more' else 'B4'].value == '\n'.join(group[side])
            assert len(list(book['赋值明细'].iter_rows(min_row=2))) == 3
        finally:
            book.close()
        assert [path.read_bytes() for path in paths] == original_files
    finally:
        bridge.shutdown()
