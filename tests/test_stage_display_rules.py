from pathlib import Path
import json
import re

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import stage_one
from engine.differ import assemble_backup, diff
from engine.ladder import build_ladder
from engine.mapper import map_raw_to_ladder
from engine.models import Cell, RawRow, RawTable, Trim, ValuationItem, ValuationTable
from engine.rules import Rules
from engine.valuer import default_valuation


PLAN = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]


def effective_exclusions(*extra):
    return list(dict.fromkeys((*stage_one.DEFAULT_EXCLUDED_ROWS, *extra)))


def raw_table(rows):
    return RawTable(model='全量展示测试车', source='rich', trims=[
        Trim(0, '基础', price_guide=10), Trim(1, '升级', price_guide=11),
    ], rows=rows)


def items(raw, index=1, base=0, excluded_rows=()):
    features = stage_one.features(raw, excluded_rows=excluded_rows)
    return stage_one._compact_items(stage_one._change_items(features, index, base), base)


def test_unknown_standard_presence_shows_name_and_all_values_are_compared():
    raw = raw_table([
        RawRow('智能香氛', [Cell(), Cell(dot='●')]),
        RawRow('屏幕刷新率(Hz)', [Cell(text='60'), Cell(text='120')]),
        RawRow('光感遮阳帘', [Cell(dot='●'), Cell()]),
    ])
    result = items(raw)
    assert '智能香氛' in result
    assert not {'有', '支持', '是', '●'}.intersection(result)
    assert any('屏幕刷新率(Hz)' in value and '120' in value and '60' in value
               and re.search(r'[（(].*60.*[）)]$', value) for value in result)
    assert '无（光感遮阳帘）' in result


def test_optional_unknown_equipment_is_separate_and_exclusion_hides_both_markers():
    raw = raw_table([
        RawRow('智能香氛', [Cell(), Cell(dot='●', subs=[Cell(dot='○', text='香氛盒(300元)')])]),
        RawRow('自动遮阳帘', [Cell(), Cell(dot='○')]),
    ])
    assert items(raw) == ['智能香氛']
    assert stage_one._option_lines(raw, 1) == ['智能香氛：香氛盒(300元)', '自动遮阳帘']
    hidden = ['智能香氛', '自动遮阳帘']
    assert items(raw, excluded_rows=hidden) == []
    assert stage_one._option_lines(raw, 1, excluded_rows=hidden) == []
    markdown = stage_one.render(raw, PLAN)
    standard, optional = markdown.split('## 选装')
    assert '自动遮阳帘' not in standard and '香氛盒' not in standard
    assert '自动遮阳帘' in optional and '香氛盒' in optional
    assert '智能香氛' not in stage_one.render(raw, PLAN, excluded_rows=hidden)


def test_known_seat_row_keeps_unrecognized_standard_subfunctions():
    raw = raw_table([RawRow('前排座椅功能', [Cell(), Cell(subs=[
        Cell(dot='●', text='加热'), Cell(dot='●', text='腿托'),
        Cell(dot='○', text='座椅旋转(2000元)'),
    ])])])
    result = items(raw)
    assert any('加热' in value for value in result)
    assert any('腿托' in value for value in result)
    assert not any('座椅旋转' in value for value in result)
    assert any('座椅旋转' in value for value in stage_one._option_lines(raw, 1, excluded_rows=[]))
    assert items(raw, excluded_rows=['前排座椅功能']) == []


def test_duplicate_unknown_names_keep_both_groups():
    raw = raw_table([
        RawRow('遮阳帘', [Cell(), Cell(text='电动')], group='天窗/玻璃'),
        RawRow('遮阳帘', [Cell(), Cell(text='手动')], group='车窗'),
    ])
    result = items(raw)
    assert any('电动' in value for value in result)
    assert any('手动' in value for value in result)
    assert len(stage_one.features(raw)) == 2
    assert items(raw, excluded_rows=['遮阳帘']) == []


def test_duplicate_unknown_rows_without_groups_are_not_silently_discarded():
    raw = raw_table([
        RawRow('遮阳帘', [Cell(), Cell(text='电动')]),
        RawRow('遮阳帘', [Cell(), Cell(text='手动')]),
    ])
    result = items(raw)
    assert any('电动' in value for value in result)
    assert any('手动' in value for value in result)
    assert len(stage_one.features(raw)) == 2
    assert items(raw, excluded_rows=['遮阳帘']) == []


@pytest.mark.parametrize('hidden, expected, optional', [
    ('手机无线充电功率', '前排手机无线充电', []),
    ('手机无线充电功能', '50W手机无线充电', ['手机无线充电功率：80W']),
])
def test_excluded_sources_cannot_reappear_through_merged_feature_lookups(hidden, expected, optional):
    raw = raw_table([
        RawRow('手机无线充电功能', [Cell(), Cell(dot='●', text='前排')]),
        RawRow('手机无线充电功率', [Cell(), Cell(dot='●', text='50W', subs=[Cell(dot='○', text='80W')])]),
    ])
    original = raw.to_json()
    assert items(raw) == ['前排50W手机无线充电']
    assert items(raw, excluded_rows=[hidden]) == [expected]
    assert stage_one._option_lines(raw, 1, excluded_rows=[hidden]) == optional
    assert raw.to_json() == original


def test_confirmed_exclusions_are_defaults_but_empty_list_reveals_raw_rows():
    raw = raw_table([
        RawRow('电池容量(kWh)', [Cell(text='40'), Cell(text='50')]),
        RawRow('电池能量(kWh)', [Cell(text='40'), Cell(text='50')]),
        RawRow('备胎规格', [Cell(), Cell(text='非全尺寸')]),
        RawRow('厂商指导价(元)', [Cell(text='100000'), Cell(text='110000')]),
        RawRow('智能香氛', [Cell(), Cell(dot='●')]),
    ])
    default_keys = {feature['key'] for feature in stage_one.features(raw)}
    assert '电池容量(kWh)' not in default_keys and '备胎规格' not in default_keys
    assert {'电池能量(kWh)', '智能香氛'}.issubset(default_keys)
    all_keys = {feature['key'] for feature in stage_one.features(raw, excluded_rows=[])}
    assert {row.name for row in raw.rows}.issubset(all_keys)


def test_all_confirmed_defaults_hide_standard_and_optional_but_keep_raw_data():
    defaults = effective_exclusions()
    assert {
        '三电系统质保', '高精定位', '记忆泊车', '循迹倒车', '信号灯识别',
        '隐藏电动门把手', '主动降噪', 'USB/Type-C最大充电功率',
        '对外直流放电功率(kW)', '陡坡缓降', '起步提醒', '语音免唤醒词', '面部识别', '激光雷达线数',
        '系统综合马力(Ps)',
    }.issubset(defaults)
    assert not {'樱珞粉主题套装', '牛油果绿主题套装', '罗兰紫主题套装'}.intersection(defaults)
    raw = raw_table([RawRow(name, [Cell(dot='●'), Cell(dot='○')]) for name in defaults] + [
        RawRow('新发现的配置', [Cell(), Cell(dot='●', subs=[Cell(dot='○', text='扩展包')])]),
    ])
    original = raw.to_json()
    assert [row.name for row in stage_one.display_raw(raw).rows] == ['新发现的配置']
    assert items(raw, excluded_rows=None) == ['新发现的配置']
    assert stage_one._option_lines(raw, 1) == ['新发现的配置：扩展包']
    assert stage_one.display_raw(raw, excluded_rows=[]).rows == raw.rows
    assert raw.to_json() == original


def test_new_exclusions_match_exact_names_without_hiding_other_usb_rows_or_theme_packs():
    visible = ['USB/Type-C接口数量', '其他主题套装', '三电系统质保扩展政策']
    raw = raw_table([RawRow(name, [Cell(), Cell(dot='●')]) for name in visible])
    assert [row.name for row in stage_one.display_raw(raw).rows] == visible


def test_lidar_line_count_is_hidden_by_default_but_restorable_without_valuation_changes():
    raw = raw_table([
        RawRow('激光雷达数量', [Cell(dot='●', text='1'), Cell(dot='●', text='1')]),
        RawRow('激光雷达品牌', [Cell(dot='●', text='华为'), Cell(dot='●', text='华为')]),
        RawRow('激光雷达线数', [Cell(dot='●', text='192'), Cell(dot='●', text='896')]),
    ])
    original = raw.to_json()
    assert '激光雷达线数' not in {feature['key'] for feature in stage_one.features(raw)}
    assert '激光雷达线数' in {feature['key'] for feature in stage_one.features(raw, excluded_rows=[])}
    assert stage_one.display_raw(raw, excluded_rows=[]).rows == raw.rows
    rules = Rules()
    mapped = map_raw_to_ladder(raw, rules)
    assert mapped[8]['values'][0] == mapped[8]['values'][1]
    assert map_raw_to_ladder(stage_one.display_raw(raw), rules) == mapped
    assert raw.to_json() == original


@pytest.mark.parametrize('restore', [False, True])
def test_nedc_combined_range_is_hidden_in_all_exports_and_can_be_restored_after_restart(tmp_path, restore):
    target, fuel, pure = 'NEDC综合续航(km)', 'NEDC综合油耗(L/100km)', 'NEDC纯电续航里程(km)'
    raw = raw_table([
        RawRow(target, [Cell(text='1000'), Cell(text='1200')], group='电池/充电'),
        RawRow(fuel, [Cell(text='5.0'), Cell(text='4.5')], group='基本参数'),
        RawRow(pure, [Cell(text='100'), Cell(text='150')], group='电池/充电'),
    ])
    original = raw.to_json()
    rules = Rules()
    mapped = map_raw_to_ladder(raw, rules)
    assert target not in stage_one.RANGE_ROWS
    keys = {feature['key'] for feature in stage_one.features(raw)}
    assert target not in keys and fuel not in keys and pure in keys
    workdir = tmp_path / 'work'
    bridge = Bridge(str(workdir))
    bridge.stage_raw = raw
    try:
        settings = bridge.stage_display_rules()
        assert settings['ok']
        assert target in settings['default_excluded_rows'] and target in settings['excluded_rows']
        if restore:
            saved = bridge.save_stage_display_rules([name for name in settings['excluded_rows'] if name != target])
            assert saved['ok'] and target not in saved['excluded_rows']
        preview = bridge.stage_preview(PLAN)
        assert preview['ok']
        assert any(target in item for column in preview['columns'] for item in column['items']) == restore
    finally:
        bridge.shutdown()
    restarted = Bridge(str(workdir))
    restarted.stage_raw = raw
    restarted.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        settings = restarted.stage_display_rules()
        assert target in settings['default_excluded_rows']
        assert (target not in settings['excluded_rows']) == restore
        preview = restarted.stage_preview(PLAN)
        markdown = restarted.stage_export(PLAN)
        excel = restarted.export_excel('ladder', PLAN)
        assert preview['ok'] and markdown['ok'] and excel['ok']
        preview_text = '\n'.join(value for column in preview['columns'] for value in [*column['items'], *column['options']])
        book = load_workbook(excel['path'])
        try:
            excel_text = '\n'.join(str(cell.value or '') for row in book['配置阶梯'] for cell in row)
        finally:
            book.close()
        for text in [preview_text, markdown['md'], excel_text]:
            assert (target in text) == restore
            assert fuel not in text
            assert 'NEDC纯电续航' in text
        assert raw.to_json() == original
        assert map_raw_to_ladder(raw, rules) == mapped
    finally:
        restarted.shutdown()


def test_saved_empty_list_restores_all_defaults_and_survives_restart(tmp_path):
    bridge = Bridge(str(tmp_path))
    try:
        rules = bridge.stage_display_rules()
        assert rules['excluded_rows'] == effective_exclusions()
        assert rules['default_excluded_rows'] == effective_exclusions()
        saved = bridge.save_stage_display_rules([])
        assert saved['ok'] and saved['excluded_rows'] == []
        data = json.loads((tmp_path / '阶梯展示排除清单.json').read_text(encoding='utf-8'))
        assert data == {'schema': 'carkit.stage-display/v5', 'excluded_rows': [],
                        'included_rows': effective_exclusions()}
    finally:
        bridge.shutdown()
    restarted = Bridge(str(tmp_path))
    try:
        rules = restarted.stage_display_rules()
        assert rules['excluded_rows'] == []
        assert rules['default_excluded_rows'] == effective_exclusions()
    finally:
        restarted.shutdown()


@pytest.mark.parametrize('legacy', [
    ['个人隐藏项', '三电系统质保', '个人隐藏项'],
    {'schema': 'carkit.stage-display/v2', 'excluded_rows': ['个人隐藏项'], 'defaults_applied': []},
    {'schema': 'carkit.stage-display/v3', 'excluded_rows': ['个人隐藏项'],
     'defaults_applied': ['三电系统质保']},
    {'excluded_rows': ['个人隐藏项', '三电系统质保']},
])
def test_old_lists_immediately_apply_all_code_exclusions_and_keep_only_extras(tmp_path, monkeypatch, legacy):
    path = tmp_path / '阶梯展示排除清单.json'
    path.write_text(json.dumps(legacy, ensure_ascii=False), encoding='utf-8')
    # A code addition works even before the old directory has ever been upgraded.
    added = '升级前新增默认隐藏配置'
    monkeypatch.setattr(stage_one, 'DEFAULT_EXCLUDED_ROWS', (*stage_one.DEFAULT_EXCLUDED_ROWS, added))
    bridge = Bridge(str(tmp_path))
    try:
        bridge.stage_raw = raw_table([
            RawRow(name, [Cell(), Cell(dot='●')])
            for name in ['个人隐藏项', '三电系统质保', added, '保持展示项']
        ])
        rules = bridge.stage_display_rules()
        assert rules['ok'] and rules['excluded_rows'] == effective_exclusions('个人隐藏项')
        assert rules['default_excluded_rows'] == effective_exclusions()
        assert bridge.stage_preview(PLAN)['columns'][1]['items'] == ['保持展示项']
        saved = json.loads(path.read_text(encoding='utf-8'))
        assert saved == {'schema': 'carkit.stage-display/v5', 'excluded_rows': ['个人隐藏项'],
                         'included_rows': []}
        migrated = path.read_bytes(), path.stat().st_mtime_ns
        assert bridge.stage_display_rules() == rules
        assert (path.read_bytes(), path.stat().st_mtime_ns) == migrated
        # Clearing only the extras keeps the initially selected defaults hidden.
        assert bridge.save_stage_display_rules(effective_exclusions())['excluded_rows'] == effective_exclusions()
        assert bridge.stage_preview(PLAN)['columns'][1]['items'] == ['个人隐藏项', '保持展示项']
    finally:
        bridge.shutdown()


def test_v4_lists_keep_defaults_until_user_explicitly_restores_them(tmp_path):
    path = tmp_path / '阶梯展示排除清单.json'
    path.write_text(json.dumps({'schema': 'carkit.stage-display/v4', 'excluded_rows': ['个人隐藏项']}),
                    encoding='utf-8')
    original = path.read_bytes()
    bridge = Bridge(str(tmp_path))
    try:
        assert bridge.stage_display_rules()['excluded_rows'] == effective_exclusions('个人隐藏项')
        assert path.read_bytes() == original
        restored = '三电系统质保'
        selected = [name for name in effective_exclusions('个人隐藏项') if name != restored]
        assert bridge.save_stage_display_rules(selected)['excluded_rows'] == selected
        assert json.loads(path.read_text(encoding='utf-8')) == {
            'schema': 'carkit.stage-display/v5', 'excluded_rows': ['个人隐藏项'],
            'included_rows': [restored],
        }
    finally:
        bridge.shutdown()


def test_restored_default_survives_restart_new_defaults_and_can_be_hidden_again(tmp_path, monkeypatch):
    restored, added = '三电系统质保', '后续新增默认隐藏配置'
    raw = raw_table([RawRow(name, [Cell(), Cell(dot='●')]) for name in [restored, added]])
    selected = [name for name in effective_exclusions('个人隐藏项') if name != restored]
    bridge = Bridge(str(tmp_path))
    try:
        bridge.stage_raw = raw
        assert bridge.save_stage_display_rules(selected)['excluded_rows'] == selected
        assert bridge.stage_preview(PLAN)['columns'][1]['items'] == [restored, added]
    finally:
        bridge.shutdown()
    monkeypatch.setattr(stage_one, 'DEFAULT_EXCLUDED_ROWS', (*stage_one.DEFAULT_EXCLUDED_ROWS, added, added))
    restarted = Bridge(str(tmp_path))
    try:
        restarted.stage_raw = raw
        expected = [name for name in effective_exclusions('个人隐藏项') if name != restored]
        assert restarted.stage_display_rules()['excluded_rows'] == expected
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == [restored]
        saved = restarted.save_stage_display_rules([*expected, restored, restored])
        assert saved['excluded_rows'] == effective_exclusions('个人隐藏项')
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == []
        data = json.loads((tmp_path / '阶梯展示排除清单.json').read_text(encoding='utf-8'))
        assert data['included_rows'] == []
    finally:
        restarted.shutdown()


def test_code_exclusions_apply_additions_deletions_and_duplicates_without_sticky_saved_defaults(tmp_path, monkeypatch):
    defaults = stage_one.DEFAULT_EXCLUDED_ROWS
    first, second = '后续新增默认隐藏配置', '再次新增默认隐藏配置'
    hidden = ['个人隐藏项']
    raw = raw_table([
        RawRow(name, [Cell(), Cell(dot='●')])
        for name in [*hidden, first, second, '保持展示项']
    ])
    path = tmp_path / '阶梯展示排除清单.json'
    bridge = Bridge(str(tmp_path))
    try:
        bridge.stage_raw = raw
        assert bridge.save_stage_display_rules(effective_exclusions(*hidden))['ok']
        monkeypatch.setattr(stage_one, 'DEFAULT_EXCLUDED_ROWS', (*defaults, first, first))
        rules = bridge.stage_display_rules()
        assert rules['excluded_rows'] == effective_exclusions(*hidden)
        assert rules['excluded_rows'].count(first) == 1
        assert rules['default_excluded_rows'].count(first) == 1
        assert bridge.stage_preview(PLAN)['columns'][1]['items'] == [second, '保持展示项']
        before = path.read_bytes(), path.stat().st_mtime_ns
        assert bridge.stage_display_rules()['excluded_rows'] == effective_exclusions(*hidden)
        assert (path.read_bytes(), path.stat().st_mtime_ns) == before
        saved = bridge.save_stage_display_rules([*effective_exclusions(*hidden), first, ' 个人隐藏项 '])
        assert saved['ok'] and saved['excluded_rows'] == effective_exclusions(*hidden)
        data = json.loads(path.read_text(encoding='utf-8'))
        assert data == {'schema': 'carkit.stage-display/v5', 'excluded_rows': hidden, 'included_rows': []}
    finally:
        bridge.shutdown()

    restarted = Bridge(str(tmp_path))
    try:
        restarted.stage_raw = raw
        assert restarted.stage_display_rules()['excluded_rows'] == effective_exclusions(*hidden)
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == [second, '保持展示项']
        monkeypatch.setattr(stage_one, 'DEFAULT_EXCLUDED_ROWS', (second, *defaults, first))
        assert restarted.stage_display_rules()['excluded_rows'] == effective_exclusions(*hidden)
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == ['保持展示项']
        monkeypatch.setattr(stage_one, 'DEFAULT_EXCLUDED_ROWS', (*defaults, second))
        assert first not in restarted.stage_display_rules()['excluded_rows']
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == [first, '保持展示项']
        assert json.loads(path.read_text(encoding='utf-8'))['excluded_rows'] == hidden
    finally:
        restarted.shutdown()

    restarted = Bridge(str(tmp_path))
    try:
        assert restarted.stage_display_rules()['excluded_rows'] == effective_exclusions(*hidden)
        assert first not in restarted.stage_display_rules()['excluded_rows']
    finally:
        restarted.shutdown()


def test_saved_exclusions_apply_to_history_survive_restart_and_can_be_restored(tmp_path):
    workdir = tmp_path / 'work'
    bridge = Bridge(str(workdir))
    capture = workdir / 'raw' / 'capture.json'
    raw_table([
        RawRow('智能香氛', [Cell(), Cell(dot='●')]),
        RawRow('智能香氛', [Cell(), Cell(dot='○', text='扩展包')], group='特色配置'),
        RawRow('自动遮阳帘', [Cell(), Cell(dot='●')]),
    ]).save(capture)
    original = capture.read_bytes()
    try:
        assert bridge.open_history(capture.name)['ok']
        initial = bridge.stage_display_rules()
        assert initial['ok'] and initial['rows'] == ['智能香氛', '自动遮阳帘']
        assert bridge.save_stage_display_rules(effective_exclusions(' 智能香氛 ', '智能香氛'))['ok']
        assert bridge.stage_display_rules()['excluded_rows'] == effective_exclusions('智能香氛')
        preview = bridge.stage_preview(PLAN)
        assert preview['ok'] and preview['columns'][1]['items'] == ['自动遮阳帘']
        assert preview['columns'][1]['options'] == []
    finally:
        bridge.shutdown()
    restarted = Bridge(str(workdir))
    try:
        assert restarted.stage_display_rules()['excluded_rows'] == effective_exclusions('智能香氛')
        assert restarted.open_history(capture.name)['ok']
        assert restarted.stage_preview(PLAN)['columns'][1]['items'] == ['自动遮阳帘']
        assert restarted.save_stage_display_rules(effective_exclusions())['ok']
        assert '智能香氛' in restarted.stage_preview(PLAN)['columns'][1]['items']
        assert restarted.stage_preview(PLAN)['columns'][1]['options'] == ['智能香氛：扩展包']
        assert capture.read_bytes() == original
        # Display preferences belong to the selected data folder, not a global setting.
        other = Bridge(str(tmp_path / 'other'))
        try:
            assert '智能香氛' not in other.stage_display_rules()['excluded_rows']
        finally:
            other.shutdown()
    finally:
        restarted.shutdown()


@pytest.mark.parametrize('invalid', ['智能香氛', {'excluded_rows': []}, None, [3], ['有效配置', None]])
def test_invalid_exclusion_input_preserves_last_saved_configuration(tmp_path, invalid):
    bridge = Bridge(str(tmp_path))
    try:
        assert bridge.save_stage_display_rules(effective_exclusions('智能香氛'))['ok']
        saved = {path.relative_to(tmp_path): path.read_bytes() for path in tmp_path.rglob('*.json')}
        result = bridge.save_stage_display_rules(invalid)
        assert result['ok'] is False and result.get('error')
        assert bridge.stage_display_rules()['excluded_rows'] == effective_exclusions('智能香氛')
        assert {path.relative_to(tmp_path): path.read_bytes() for path in tmp_path.rglob('*.json')} == saved
    finally:
        bridge.shutdown()


@pytest.mark.parametrize('restore_default', [False, True])
def test_preview_markdown_excel_share_full_data_and_saved_exclusions(tmp_path, restore_default):
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw_table([
        RawRow('智能香氛', [Cell(), Cell(dot='●')]),
        RawRow('屏幕刷新率(Hz)', [Cell(text='60'), Cell(text='120')]),
        RawRow('光感遮阳帘', [Cell(dot='●'), Cell()]),
        RawRow('自动遮阳帘', [Cell(), Cell(dot='○')]),
        RawRow('隐藏的额外功能', [Cell(), Cell(dot='●')]),
        RawRow('三电系统质保', [Cell(), Cell(dot='●', text='终身', subs=[Cell(dot='○', text='延保包')])]),
    ])
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        excluded = effective_exclusions('隐藏的额外功能')
        if restore_default:
            excluded.remove('三电系统质保')
        assert bridge.save_stage_display_rules(excluded)['ok']
        preview = bridge.stage_preview(PLAN)
        markdown = bridge.stage_export(PLAN, choose_path=True)
        excel = bridge.export_excel('ladder', PLAN)
        assert preview['ok'] and markdown['ok'] and excel['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        hidden = ['隐藏的额外功能'] + ([] if restore_default else ['三电系统质保', '延保包'])
        assert not any(name in markdown['md'] for name in hidden)
        assert any('三电系统质保' in item for item in preview['columns'][1]['items']) == restore_default
        assert ('三电系统质保：延保包' in preview['columns'][1]['options']) == restore_default
        for column in preview['columns']:
            assert not any(name in value for name in hidden
                           for value in [*column['items'], *column['options']])
            for value in column['items']:
                assert value in markdown['md']
            for value in column['options']:
                assert '> 选装：' + value in markdown['md']
        book = load_workbook(excel['path'])
        try:
            sheet = book['配置阶梯']
            for row, column in enumerate(preview['columns'], 2):
                assert (sheet.cell(row, 5).value or '') == '\n'.join(column['items'])
                assert (sheet.cell(row, 6).value or '') == '\n'.join(column['options'])
        finally:
            book.close()
    finally:
        bridge.shutdown()


def test_display_exclusions_do_not_change_mapping_or_competitive_valuation(tmp_path):
    bridge = Bridge(str(tmp_path))
    raw = raw_table([
        RawRow('AR-HUD增强现实抬头显示', [Cell(), Cell(dot='●')]),
        RawRow('智能香氛', [Cell(), Cell(dot='●')]),
    ])
    bridge.stage_raw = raw
    source = raw.to_json()
    def competition():
        rules = Rules()
        ladder = build_ladder(raw, rules)
        valuation = ValuationTable(items=[ValuationItem(**item) for item in default_valuation(rules)['items']])
        pairs = [{'self_trim': '升级', 'comp_trim': '基础'}]
        cells = diff(ladder, ladder, pairs, rules, valuation=valuation)
        groups = assemble_backup(cells, pairs, raw.model, raw.model, {'升级': 11}, {'基础': 10}, valuation)
        return map_raw_to_ladder(raw, rules), cells, groups
    try:
        before = competition()
        assert before[2][0]['valuation']['config_adv'] == 2000
        assert bridge.save_stage_display_rules(effective_exclusions('AR-HUD增强现实抬头显示', '智能香氛'))['ok']
        assert bridge.stage_preview(PLAN)['columns'][1]['items'] == []
        assert competition() == before
        assert raw.to_json() == source
    finally:
        bridge.shutdown()
