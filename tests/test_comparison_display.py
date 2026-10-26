"""Competition output shares the ladder wording without changing valuation."""
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import differ, render_backup, valuer
from engine.models import Ladder, LadderItem, ValuationItem, ValuationTable
from engine.rules import Rules, display_order_key


RULES = Rules()
VALUATION = ValuationTable(items=[ValuationItem(**item)
                                 for item in valuer.default_valuation(RULES)['items']])
PAIR = {'self_trim': '左版', 'comp_trim': '右版',
        'self_floor_price': 100000, 'comp_floor_price': 102000}


def ladder(values, name, model):
    return Ladder(model=model, trims=[{'name': name, 'price_guide': 10}], items=[
        LadderItem(no=item['no'], name=item['name'], values=[values.get(item['no'], '✕')])
        for item in RULES.items])


def result(left, right):
    cells = differ.diff(ladder(left, '左版', '本品'), ladder(right, '右版', '竞品'),
                        [PAIR], RULES, valuation=VALUATION)
    groups = differ.assemble_backup(cells, [PAIR], '本品', '竞品',
                                    {'左版': 10}, {'右版': 10}, VALUATION)
    return cells, groups


@pytest.mark.parametrize('no,value,expected', [
    (3, '●800V', '800V'),
    (18, '●不可开启全景(光感天幕)', '光感天幕'),
    (18, '●光感天幕(分区调光)', '光感天幕(分区调光)'),
    (18, '●不可开启全景(光感天幕+电致变色)', '光感天幕+电致变色'),
    (18, '●不可开启全景', '不可开启全景天窗'),
    (20, '外后视镜电动调节+外后视镜电动折叠+锁车自动折叠+外后视镜加热+记忆+倒车自动下翻',
     '外后视镜电调+折叠+加热+记忆+倒车下翻'),
    (21, '15.6寸中控', '15.6寸中控'),
    (22, '12.3寸副驾屏', '12.3寸副驾屏'),
    (29, '全液晶仪表(8.8寸)', '全液晶8.8寸仪表'),
    (29, '10.17寸仪表', '10.17寸仪表'),
    (29, '全液晶仪表', '[待定]全液晶仪表'),
    (34, '仿皮', '仿皮座椅'),
    (35, '主6副4电调', '主驾6向电调+副驾4向电调'),
    (35, '主副手调', '主驾6向手调+副驾4向手调'),
    (35, '主驾电调+副驾手调', '主驾电调+副驾手调'),
])
def test_competition_configuration_values_use_requested_wording(no, value, expected):
    cells, groups = result({no: value}, {})
    render_backup.format_configurations(cells, groups)
    assert next(cell for cell in cells if cell['no'] == no)['self_val'] == expected


def test_formatting_only_changes_presentation_and_preserves_pair_difference_information():
    left = {3: '●800V', 18: '●不可开启全景(光感天幕)',
            20: '电调+折叠+加热+记忆+倒车下翻', 21: '15.6寸中控', 22: '12.3寸副驾屏',
            29: '全液晶仪表(8.8寸)', 34: '仿皮',
            35: '主驾6向电调+副驾4向手调(主含腰撑4向)'}
    right = {3: '●400V', 20: '电调', 21: '10寸中控', 29: '4寸仪表', 34: '织物',
             35: '主驾6向手调+副驾4向手调(主含腰撑4向)'}
    cells, groups = result(left, right)
    original_cells, original_groups = deepcopy(cells), deepcopy(groups)
    render_backup.format_configurations(cells, groups)

    for old, cell in zip(original_cells, cells):
        assert {key: value for key, value in cell.items()
                if key not in ('self_val', 'comp_val', 'display', 'backup_more', 'backup_less')} == {
                    key: value for key, value in old.items()
                    if key not in ('self_val', 'comp_val', 'display', 'backup_more', 'backup_less')}
    for old, group in zip(original_groups, groups):
        assert {key: value for key, value in group.items()
                if key not in ('more', 'less', 'valuation')} == {
                    key: value for key, value in old.items()
                    if key not in ('more', 'less', 'valuation')}
        assert {key: value for key, value in group['valuation'].items() if key != 'detail'} == {
            key: value for key, value in old['valuation'].items() if key != 'detail'}
        for before, after in zip(old['valuation']['detail'], group['valuation']['detail']):
            assert {key: value for key, value in after.items() if key != 'display'} == {
                key: value for key, value in before.items() if key != 'display'}
            cell = next(cell for cell in cells if cell['no'] == after['no'])
            assert after['display'] == (cell['backup_more'] or cell['backup_less'] or cell['display'])

    by_no = {cell['no']: cell for cell in cells}
    assert by_no[3]['backup_more'] == '800V'
    assert by_no[18]['backup_more'] == '光感天幕'
    assert by_no[20]['backup_more'] == '外后视镜电调+折叠+加热+记忆+倒车下翻(外后视镜电调)'
    assert by_no[29]['backup_more'] == '全液晶8.8寸仪表(4寸仪表)'
    assert '仿皮座椅' in by_no[34]['backup_more'] and '座椅调节' not in by_no[34]['backup_more']
    assert by_no[35]['backup_more'] == '主驾电调（主驾手调）'
    assert '副驾' not in by_no[35]['backup_more']  # The unchanged passenger seat is not repeated.
    assert '腰撑' not in by_no[35]['self_val'] and '腰撑' not in by_no[35]['comp_val']
    assert '腰撑' not in by_no[35]['backup_more']
    assert by_no[35]['backup_more'] in groups[0]['more']


def test_adjustment_difference_ignores_support_notes_and_preserves_direction_difference():
    cells, groups = result(
        {35: '主驾6向电调+副驾4向电调'},
        {35: '主驾6向电调+副驾6向电调(主腰撑4向/副腰撑4向)'},
    )
    original_amount = groups[0]['valuation']['config_adv']
    render_backup.format_configurations(cells, groups)
    seat = next(cell for cell in cells if cell['no'] == 35)
    assert seat['backup_less'] == '副驾4向（副驾6向）'
    assert seat['backup_less'] in groups[0]['less']
    assert groups[0]['valuation']['config_adv'] == original_amount
    assert seat['comp_config'].endswith('(主腰撑4向/副腰撑4向)')
    assert '腰撑' not in seat['display'] and '无（' not in seat['display']


def test_optional_and_pending_descriptions_do_not_become_standard_or_change_amounts():
    cells, groups = result({18: '○不可开启全景(光感天幕)(选装)',
                            29: '8.8寸仪表(○全液晶选装)',
                            35: '[待定]座椅调节：主驾方向未说明'},
                           {29: '全液晶仪表(8.8寸)', 35: '主6副4电调'})
    original_cells, original_groups = deepcopy(cells), deepcopy(groups)
    render_backup.format_configurations(cells, groups)
    by_no = {cell['no']: cell for cell in cells}
    assert by_no[18]['verdict'] == '同' and not by_no[18]['backup_more'] and not by_no[18]['backup_less']
    assert by_no[18]['self_val'] == '无' and by_no[18]['self_config'] == '✕'
    assert by_no[18]['self_source_config'] == '○不可开启全景(光感天幕)(选装)'
    assert by_no[29]['self_val'] == '8.8寸仪表'
    assert by_no[29]['self_config'] == '8.8寸仪表'
    assert by_no[29]['self_source_config'] == '8.8寸仪表(○全液晶选装)'
    assert by_no[29]['verdict'] == '少'
    assert not any(marker in by_no[29]['display'] for marker in ('○', '选装', '选配'))
    assert '[待定]' in by_no[35]['self_val'] and by_no[35]['verdict'] == '不计'
    assert not by_no[35]['backup_more'] and not by_no[35]['backup_less']
    assert [cell['verdict'] for cell in cells] == [cell['verdict'] for cell in original_cells]
    assert {key: value for key, value in groups[0]['valuation'].items() if key != 'detail'} == {
        key: value for key, value in original_groups[0]['valuation'].items() if key != 'detail'}


def test_api_markdown_and_excel_share_configuration_display_without_changing_saved_ladders(tmp_path):
    left = ladder({3: '●800V', 18: '●不可开启全景(光感天幕)', 20: '电调+折叠+加热',
                   21: '15.6寸中控', 29: '全液晶仪表(8.8寸)', 34: '仿皮',
                   35: '主驾6向电调+副驾4向手调'}, '左版', '本品')
    right = ladder({3: '●400V', 20: '电调', 21: '10寸中控', 29: '4寸仪表', 34: '织物',
                    35: '主驾6向手调+副驾4向手调'}, '右版', '竞品')
    left_path, right_path = tmp_path / 'left.json', tmp_path / 'right.json'
    left.save(left_path)
    right.save(right_path)
    original_files = [path.read_bytes() for path in (left_path, right_path)]
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        report = bridge.run_diff(str(left_path), str(right_path), [PAIR], left_is_competitor=True)
        assert report['ok'], report.get('error')
        group = report['groups'][0]
        for label in ('800V', '光感天幕', '全液晶8.8寸仪表(4寸仪表)'):
            assert label in group['more']
        assert Path(report['md_path']).read_text(encoding='utf-8') == report['md']
        assert '| 本品多 | ' + '<br>'.join(group['more']) + ' |' in report['md']
        assert '| 本品少 | ' + ('<br>'.join(group['less']) or '—') + ' |' in report['md']
        exported = bridge.export_excel('diff', report)
        assert exported['ok'], exported.get('error')
        book = load_workbook(exported['path'])
        try:
            assert book['对比汇总']['B3'].value == '\n'.join(group['more'])
            assert book['对比汇总']['B4'].value == ('\n'.join(group['less']) or '—')
            detail = list(book['赋值明细'].iter_rows(min_row=2, values_only=True))
            assert [row[3] for row in detail] == [item['display'] for item in sorted(
                group['valuation']['detail'], key=lambda item: display_order_key(item['no']))]
            judgment = {row[3]: row for row in book['配置判定'].iter_rows(min_row=2, values_only=True)}
            for cell in report['cells']:
                row = judgment[cell['name']]
                assert row[4:6] == (cell['self_val'], cell['comp_val'])
                assert row[6:8] == (cell['verdict'], cell['display'])
            cluster = next(cell for cell in report['cells'] if cell['no'] == 29)
            assert cluster['self_val'] == '全液晶8.8寸仪表'
        finally:
            book.close()
        assert [path.read_bytes() for path in (left_path, right_path)] == original_files
    finally:
        bridge.shutdown()
