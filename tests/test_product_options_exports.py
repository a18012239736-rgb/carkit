import copy

from openpyxl import load_workbook

from engine.excel_export import export_report
from engine.render_backup import render_md
from test_frontend_comparison_state import run_node


def report():
    group = {
        'pair': {'self_trim': '舒适型', 'comp_trim': 'MAX'},
        'self_price': 10, 'comp_price': 11, 'more': ['800V'], 'less': [],
        'self_floor_price': 90000, 'comp_floor_price': 100000,
        'valuation': {'config_adv': 2000, 'flat_adv': 12000, 'overall': 12000,
                      'total_more': 2000, 'total_less': 0, 'detail': []},
    }
    return {'self_model': '本品车', 'comp_model': '竞品车', 'groups': [group], 'cells': []}


def markdown(data):
    return render_md(data['self_model'], data['comp_model'], data['groups'], data['cells'], date='2026-10-15')


def test_options_are_separate_from_metrics_and_do_not_mutate_report(tmp_path):
    data = report()
    data['groups'][0].update(self_options=['舒适选装包', '座椅通风'], comp_options=['方向盘加热'])
    second = copy.deepcopy(data['groups'][0])
    second.update(pair={'self_trim': '豪华型', 'comp_trim': 'PRO'}, self_options=[], comp_options=[])
    data['groups'].append(second)
    before = copy.deepcopy(data)
    text = markdown(data)
    assert '| 综合竞争力（元） | 12000 | 12000 |' in text
    assert '| 本品车选装（仅展示，不参与赋值） | 舒适选装包<br>座椅通风 | — |' in text
    assert '| 竞品车选装（仅展示，不参与赋值） | 方向盘加热 | — |' in text
    assert text.index('| 综合竞争力（元）') < text.index('| 本品车选装') < text.index('## 底价与价格计算依据')
    path = tmp_path / 'options.xlsx'
    export_report(path, 'diff', data)
    book = load_workbook(path)
    try:
        summary = book['对比汇总']
        assert [summary[f'B{row}'].value for row in (5, 6, 7)] == [2000, 12000, 12000]
        assert summary['A8'].value == '本品车选装（仅展示，不参与赋值）'
        assert summary['B8'].value == '舒适选装包\n座椅通风'
        assert summary['C8'].value == '—'
        assert summary['A9'].value == '竞品车选装（仅展示，不参与赋值）'
        assert summary['B9'].value == '方向盘加热'
        assert summary['B8'].alignment.wrap_text
        assert summary.row_dimensions[8].height >= 50
        assert book['赋值明细'].max_row == book['配置判定'].max_row == 1
    finally:
        book.close()
    assert data == before


def test_no_options_preserves_existing_markdown_and_excel_shape(tmp_path):
    data = report()
    original = markdown(data)
    data['groups'][0].update(self_options=[], comp_options=None)
    assert markdown(data) == original
    path = tmp_path / 'without-options.xlsx'
    export_report(path, 'diff', data)
    book = load_workbook(path)
    try:
        assert book['对比汇总'].max_row == 7
    finally:
        book.close()


def test_option_text_cannot_inject_markdown_table_html_or_excel_formulas(tmp_path):
    data = report()
    data['groups'][0]['self_options'] = ['<script>alert(1)</script>|[包](https://example.com)\n**加热**', '=1+1']
    data['groups'][0]['comp_options'] = ['=1+1']
    text = markdown(data)
    assert '<script>' not in text
    assert '&lt;script&gt;alert(1)&lt;/script&gt;&#124;\\[包\\](https://example.com)<br>\\*\\*加热\\*\\*<br>=1+1' in text
    option_row = next(line for line in text.splitlines() if line.startswith('| 本品车选装'))
    assert option_row.count('|') == 3
    path = tmp_path / 'safe-options.xlsx'
    export_report(path, 'diff', data)
    book = load_workbook(path)
    try:
        assert book['对比汇总']['B8'].data_type == 's'
        assert book['对比汇总']['B8'].value.endswith('\n=1+1')
        assert book['对比汇总']['B9'].value == '=1+1'
        assert book['对比汇总']['B9'].data_type == 's'
    finally:
        book.close()


def test_frontend_options_are_escaped_separate_rows_and_absent_for_old_results():
    run_node(r"""
load('const esc =', '/* ---------- ③ 快照');
load('function renderDiff(', '$("#btn-export-result").addEventListener');
const group={pair:{self_trim:'舒适型',comp_trim:'MAX'},self_price:10,comp_price:11,more:[],less:[],
  valuation:{config_adv:2000,flat_adv:12000,overall:12000,total_more:2000,total_less:0,detail:[]},
  self_floor_price:90000,comp_floor_price:100000};
context.renderDiff({self_model:'本品车',comp_model:'竞品车',groups:[group],cells:[]});
const old=node('#backup-table').innerHTML;
assert(!old.includes('选装'));
group.self_options=['<img src=x onerror=alert(1)>','舒适选装包\n座椅通风'];
group.comp_options=['方向盘加热'];
const result={self_model:'本品车',comp_model:'竞品车',groups:[group,{...group,self_options:[],comp_options:[]}],cells:[]};
const before=JSON.stringify(result);
context.renderDiff(result);
const html=node('#backup-table').innerHTML;
assert(html.includes('本品车选装'));assert(html.includes('竞品车选装'));
assert(html.includes('仅展示，不参与赋值'));assert(!html.includes('<img'));
assert(html.includes('&lt;img src=x onerror=alert(1)&gt;<br>舒适选装包<br>座椅通风'));
assert(html.indexOf('综合竞争力（元）')<html.indexOf('本品车选装'));
assert.strictEqual(JSON.stringify(result),before);
group.self_options=[];group.comp_options=[];
context.renderDiff({self_model:'本品车',comp_model:'竞品车',groups:[group],cells:[]});
assert.strictEqual(node('#backup-table').innerHTML,old);
""")
