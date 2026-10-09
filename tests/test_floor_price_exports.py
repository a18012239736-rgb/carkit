from openpyxl import load_workbook
import pytest

from engine.excel_export import export_report
from engine.render_backup import render_md


def report(left=80900, right=75741, overall=-2709, reason=''):
    group = {'pair': {'self_trim': '基础型', 'comp_trim': '405MAX'},
             'more': [], 'less': [], 'self_price': 8.99, 'comp_price': 8.99,
             'self_floor_price': left, 'comp_floor_price': right,
             'valuation': {'config_adv': 2450, 'flat_adv': 2450, 'overall': overall,
                           'overall_reason': reason, 'detail': []}}
    return {'self_model': 'T19NG', 'comp_model': '启源Q05', 'groups': [group], 'cells': []}


def test_floor_prices_and_full_calculation_survive_markdown_and_excel(tmp_path):
    data = report()
    md = render_md(data['self_model'], data['comp_model'], data['groups'], [])
    assert '| 综合竞争力（元） | -2709 |' in md
    assert '| 基础型 vs 405MAX | 80900 | 75741 | 2450 + 75741 − 80900 = -2709 |' in md
    assert '公式未设置' not in md and ',900' not in md
    path = tmp_path / 'report.xlsx'
    export_report(path, 'diff', data)
    with_book = load_workbook(path)
    try:
        summary = with_book['对比汇总']
        assert summary['B5'].value == 2450
        assert summary['B6'].value == 2450
        assert summary['A7'].value == '综合竞争力（元）'
        assert summary['B7'].value == -2709
        assert summary['B7'].number_format == '0.##;[Red]-0.##'
        price = with_book['价格计算']
        assert price['C2'].value == 8.99 and price['D2'].value == 8.99
        assert price['E2'].value == 80900 and price['F2'].value == 75741
        assert price['I2'].value == -2709
        assert price['J2'].value == '2450 + 75741 − 80900 = -2709'
    finally:
        with_book.close()


@pytest.mark.parametrize('left,right,reason', [
    (None, 75741, '缺少左侧底价'),
    (80900, None, '缺少右侧底价'),
    (None, None, '缺少左侧、右侧底价'),
])
def test_missing_floor_exports_reason_without_falling_back_to_guide_price(tmp_path, left, right, reason):
    data = report(left, right, None, reason)
    md = render_md(data['self_model'], data['comp_model'], data['groups'], [])
    assert '| 综合竞争力（元） | ' + reason + ' |' in md
    assert '| 拉平指导价优势（元） | 2450 |' in md
    assert '| 未填写 |' in md
    path = tmp_path / 'report.xlsx'
    export_report(path, 'diff', data)
    book = load_workbook(path)
    try:
        assert book['对比汇总']['B7'].value == reason
        assert book['价格计算']['E2'].value == left
        assert book['价格计算']['F2'].value == right
        assert book['价格计算']['I2'].value is None
        assert book['价格计算']['J2'].value == reason
    finally:
        book.close()


def test_large_positive_floor_amounts_export_without_plus_or_grouping():
    data = report(1000000, 1020000, 22450)
    md = render_md(data['self_model'], data['comp_model'], data['groups'], [])
    assert '| 综合竞争力（元） | 22450 |' in md
    assert '2450 + 1020000 − 1000000 = 22450' in md
    assert '1e+' not in md and '+22450' not in md and '1,000,000' not in md
