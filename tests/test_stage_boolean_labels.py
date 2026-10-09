from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import stage_one
from engine.models import Cell, RawRow, RawTable, Trim


PLAN = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]
LABELS = {
    '手机互联/映射': '手机互联/映射',
    '驾驶辅助影像': '影像',
    '方向盘位置调节': '方向盘位置调节',
    '天窗类型': '天窗类型',
    '可变悬架功能': '可变悬架功能',
    '空调温度控制方式': '空调温度控制方式',
}


def raw_table(rows):
    return RawTable(model='合成测试车', trims=[Trim(0, '基础', price_guide=10),
                                          Trim(1, '升级', price_guide=11)],
                    rows=[RawRow(name=name, cells=cells) for name, cells in rows.items()])


def feature_values(raw, index=1):
    return {f['key']: f['values'][index] for f in stage_one.features(raw)}


@pytest.mark.parametrize('text', ['', '有', '支持', '是', '●'])
def test_standard_presence_names_special_controls_without_bare_values(text):
    raw = raw_table({name: [Cell(), Cell(dot='●', text=text)] for name in LABELS})
    assert feature_values(raw) == {label: [label] for label in LABELS.values()}
    assert all(not values for values in feature_values(raw, 0).values())
    items = stage_one._compact_items(stage_one._change_items(stage_one.features(raw), 1, 0), 0)
    assert set(items) == set(LABELS.values())


@pytest.mark.parametrize('text', ['', '有', '支持', '是', '●'])
def test_presence_does_not_invent_scope_quantity_type_or_seat_functions(text):
    raw = raw_table({name: [Cell(), Cell(dot='●', text=text)] for name in [
        '手机无线充电功能', '手机无线充电功率', '车内环境氛围灯',
        '方向盘材质', '座椅材质', '轮圈材质', '激光雷达数量', '激光雷达品牌',
        '电动座椅记忆', '前排座椅功能', '第二排座椅功能', '后排座椅功能',
        '外后视镜功能', '车窗一键升降功能', '无钥匙进入功能',
    ]})
    assert feature_values(raw) == {
        '手机无线充电': ['手机无线充电'], '氛围灯': ['氛围灯'],
        '方向盘': ['方向盘材质'], '座椅材质': ['座椅材质'], '轮毂': ['轮毂'],
        '激光雷达': ['激光雷达'], '电动座椅记忆': ['电动座椅记忆'],
        '前排座椅功能': ['前排座椅功能'], '第二排座椅功能': ['第二排座椅功能'],
        '后排座椅功能': ['后排座椅功能'], '外后视镜功能': ['外后视镜功能'],
        '车窗一键升降功能': ['车窗一键升降功能'], '无钥匙进入功能': ['无钥匙进入功能'],
    }


def test_specific_attributes_remain_specific():
    rows = {
        '手机互联/映射': '支持CarPlay/HiCar', '驾驶辅助影像': '360度全景影像',
        '方向盘位置调节': '手动上下+前后调节', '天窗类型': '不可开启全景天窗',
        '可变悬架功能': '软硬+高低调节', '空调温度控制方式': '自动空调',
        '手机无线充电功能': '前排', '手机无线充电功率': '50W',
        '车内环境氛围灯': '64色', '方向盘材质': '皮质', '座椅材质': '仿皮',
        '电动座椅记忆': '前排', '前排座椅功能': '加热(仅驾驶位)',
        '第二排座椅功能': '通风',
    }
    raw = raw_table({name: [Cell(), Cell(dot='●', text=value)] for name, value in rows.items()})
    assert feature_values(raw) == {
        '手机互联/映射': ['CarPlay/HiCar'], '影像': ['360度全景影像'],
        '方向盘位置调节': ['手动上下/前后调节'], '天窗类型': ['不可开启全景天窗'],
        '可变悬架功能': ['软硬+高低调节'], '空调温度控制方式': ['自动空调'],
        '手机无线充电': ['前排50W手机无线充电'], '氛围灯': ['64色氛围灯'],
        '方向盘': ['皮质方向盘'], '座椅材质': ['仿皮座椅'],
        '电动座椅记忆': ['主驾座椅记忆', '副驾座椅记忆'],
        '前排座椅功能': ['主驾座椅加热'], '第二排座椅功能': ['后排座椅通风'],
    }


def test_optional_presence_is_named_and_stays_out_of_standard_equipment():
    raw = raw_table({
        '手机互联/映射': [Cell(), Cell(dot='●', text='支持CarPlay', subs=[Cell(dot='○', text='HiCar')])],
        '可变悬架功能': [Cell(), Cell(dot='○')],
        '天窗类型': [Cell(), Cell(dot='○', text='可开启全景天窗')],
        '空调温度控制方式': [Cell(), Cell(dot='○', text='有')],
    })
    assert feature_values(raw) == {'手机互联/映射': ['CarPlay'], '可变悬架功能': [],
                                  '天窗类型': [], '空调温度控制方式': []}
    assert stage_one._option_lines(raw, 1) == ['手机互联/映射：HiCar', '可变悬架功能',
                                              '天窗类型：可开启全景天窗', '空调温度控制方式']
    standard, optional = stage_one.render(raw, PLAN).split('## 选装')
    assert 'CarPlay' in standard and 'HiCar' not in standard
    assert '可变悬架功能' not in standard and '可开启全景天窗' not in standard
    assert '> 选装：可变悬架功能' in optional
    assert '> 选装：天窗类型：可开启全景天窗' in optional


@pytest.mark.parametrize('rows', [{}, {name: [Cell(), Cell(text='不支持')] for name in LABELS}])
def test_missing_equipment_is_not_added(rows):
    raw = raw_table(rows)
    for index in (0, 1):
        assert not any(feature_values(raw, index).values())
        assert stage_one._change_items(stage_one.features(raw), index, None) == []
        assert stage_one._option_lines(raw, index) == []


def test_presence_deduplication_depends_on_each_trim_and_keeps_existing_short_labels():
    parameters = {'中控彩色屏幕': ('中控屏幕尺寸', '15.6英寸'),
                  '车联网': ('4G/5G网络', '5G'),
                  '行车电脑显示屏幕': ('液晶仪表尺寸', '10英寸')}
    rows = {parameter: [Cell(text=value), Cell()] for parameter, value in parameters.values()}
    rows.update({name: [Cell(dot='●'), Cell(dot='●')] for name in [
        *parameters, '前/后电动车窗', '哨兵模式/千里眼', '车窗防夹手功能', '全液晶仪表盘',
    ]})
    rows['车窗一键升降功能'] = [Cell(dot='●'), Cell()]
    raw = raw_table(rows)
    first, second = feature_values(raw, 0), feature_values(raw, 1)
    for name in parameters:
        assert first[name] == []
        assert second[name] == [name]
    for name, label in [('哨兵模式/千里眼', '哨兵模式'), ('车窗防夹手功能', '防夹手'),
                        ('全液晶仪表盘', '全液晶仪表')]:
        assert first[name] == second[name] == [label]
    features = stage_one.features(raw)
    assert '前/后电动车窗' not in dict(stage_one._change_items(features, 0, None))
    assert dict(stage_one._change_items(features, 1, None))['前/后电动车窗'] == '前/后电动车窗'


def test_preview_markdown_and_excel_share_named_presence(tmp_path):
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw_table({name: [Cell(), Cell(dot='●')] for name in LABELS} | {
        '车载冰箱': [Cell(), Cell(dot='○')],
    })
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        preview = bridge.stage_preview(PLAN)
        assert preview['ok']
        assert set(preview['columns'][1]['items']) == set(LABELS.values())
        assert preview['columns'][1]['options'] == ['车载冰箱']
        markdown = bridge.stage_export(PLAN, choose_path=True)
        assert markdown['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        for label in LABELS.values():
            assert label in markdown['md']
        assert '> 选装：车载冰箱' in markdown['md']
        excel = bridge.export_excel('ladder', PLAN)
        assert excel['ok']
        book = load_workbook(excel['path'])
        try:
            sheet = book['配置阶梯']
            for row, column in enumerate(preview['columns'], 2):
                assert (sheet.cell(row, 5).value or '') == '\n'.join(column['items'])
                assert (sheet.cell(row, 6).value or '') == '\n'.join(column['options'])
        finally:
            book.close()
    finally:
        bridge.shutdown()
