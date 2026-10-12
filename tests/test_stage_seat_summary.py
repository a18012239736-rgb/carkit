from pathlib import Path

from openpyxl import load_workbook

from api.bridge import Bridge
from engine import stage_one
from engine.mapper import map_raw_to_ladder
from engine.models import Cell, RawRow, RawTable, Trim
from engine.rules import Rules
from engine.valuer import dynamic_delta


def table(rows):
    count = len(next(iter(rows.values())))
    return RawTable(model='座椅摘要测试', trims=[Trim(i, f'版型{i}', price_guide=10+i) for i in range(count)],
                    rows=[RawRow(name=name, cells=cells) for name, cells in rows.items()])


def lines(raw, target=0, base=None, excluded=()):
    return stage_one._compact_items(stage_one._change_items(stage_one.features(raw, excluded), target, base), base)


def test_seat_modes_and_directions_share_one_summary_without_repeating_electric_or_material():
    raw = table({
        '座椅材质': [Cell(text='仿皮')],
        '主座椅调节方式': [Cell(text='前后调节靠背调节高低调节(2向)腰部支撑(4向)')],
        '副座椅调节方式': [Cell(text='前后调节靠背调节')],
        '主/副驾驶座电动调节': [Cell(text='主●/副-')],
    })
    features = {feature['key']: feature['values'] for feature in stage_one.features(raw, [])}
    assert '座椅电调' not in features
    assert features['主驾座椅调节'] == [['主驾6向电调']]
    assert features['副驾座椅调节'] == [['副驾4向手调']]
    result = lines(raw)
    assert '仿皮座椅' in result
    assert '主驾6向电调，副驾4向手调' in result
    assert '主驾腰部支撑(4向)' in result
    assert not any(value in result for value in ['前排电调', '主驾电调', '主驾座椅电调'])
    raw.row('主座椅调节方式').cells[0] = Cell(text='6向电动(腰部支撑(4向))')
    result = lines(raw)
    assert '主驾6向电调，副驾4向手调' in result
    assert any('腰部支撑(4向)' in value for value in result)
    assert not any('6向电动' in value for value in result)


def test_unknown_direction_mode_and_optional_sources_are_not_filled_with_guesses():
    electric_only = table({'主/副驾驶座电动调节': [Cell(text='主●/副-')]})
    assert lines(electric_only) == ['主驾电调']
    unknown_mode = table({'主/副驾驶座电动调节': [Cell(text='[待定]')],
                          '主座椅调节方式': [Cell(text='6向')], '副座椅调节方式': [Cell(text='4向')]})
    assert lines(unknown_mode) == ['主驾6向调节，副驾4向调节']
    optional = table({'主/副驾驶座电动调节': [Cell(dot='○', text='主/副')],
                      '主座椅调节方式': [Cell(dot='○', text='6向电动')],
                      '副座椅调节方式': [Cell(dot='○', text='4向手动')]})
    assert lines(optional) == []
    assert stage_one._option_lines(optional, 0, []) == [
        '主/副驾驶座电动调节：主/副', '主座椅调节方式：6向电动', '副座椅调节方式：4向手动',
    ]


def test_hiding_seat_source_does_not_leak_direction_or_infer_mode_from_hidden_header():
    raw = table({'主/副驾驶座电动调节': [Cell(text='主●/副-')],
                 '主座椅调节方式': [Cell(text='6向')], '副座椅调节方式': [Cell(text='4向')]})
    assert lines(raw, excluded=['主/副驾驶座电动调节']) == ['主驾6向调节，副驾4向调节']
    assert lines(raw, excluded=['主座椅调节方式']) == ['主驾电调，副驾4向手调']
    assert lines(raw, excluded=[row.name for row in raw.rows]) == []
    raw.row('主座椅调节方式').cells[0] = Cell(text='6向电动')
    features = {feature['key']: feature['values'] for feature in stage_one.features(raw, ['主/副驾驶座电动调节'])}
    assert features['主驾座椅调节'] == [['主驾6向电调']]
    assert features['副驾座椅调节'] == [['副驾4向调节']]
    assert lines(raw, excluded=['主/副驾驶座电动调节']) == ['主驾6向电调，副驾4向调节']
    raw.row('主座椅调节方式').cells[0] = Cell(text='前后调节靠背调节高低调节(2向)电动腰托')
    features = {feature['key']: feature['values'] for feature in stage_one.features(raw, ['主/副驾驶座电动调节'])}
    assert features['主驾座椅调节'] == [['主驾6向调节']]
    assert lines(raw, excluded=['主/副驾驶座电动调节']) == ['主驾6向调节，副驾4向调节', '主驾电动腰托']


def test_only_changed_seat_is_shown_and_display_does_not_change_raw_mapper_or_valuation():
    raw = table({'主/副驾驶座电动调节': [Cell(text='主●/副-')] * 2,
                 '主座椅调节方式': [Cell(text='6向'), Cell(text='8向')],
                 '副座椅调节方式': [Cell(text='4向')] * 2})
    original = raw.to_json()
    mapped = map_raw_to_ladder(raw, Rules())
    amount = dynamic_delta(35, mapped[35]['values'][1], mapped[35]['values'][0])
    assert lines(raw, 1, 0) == ['主驾8向电调（主驾6向电调）']
    assert raw.to_json() == original
    after = map_raw_to_ladder(raw, Rules())
    assert after == mapped
    assert dynamic_delta(35, after[35]['values'][1], after[35]['values'][0]) == amount


def test_preview_markdown_and_excel_share_seat_summary_without_changing_material_display(tmp_path):
    raw = table({'座椅材质': [Cell(text='仿皮')] * 2,
                 '主/副驾驶座电动调节': [Cell(text='主●/副-')] * 2,
                 '主座椅调节方式': [Cell(text='6向'), Cell(text='8向')],
                 '副座椅调节方式': [Cell(text='4向')] * 2})
    plan = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]
    expected = [['仿皮座椅', '主驾6向电调，副驾4向手调'], ['主驾8向电调（主驾6向电调）']]
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        assert bridge.save_stage_display_rules([])['ok']
        preview = bridge.stage_preview(plan)
        assert preview['ok']
        assert [column['items'] for column in preview['columns']] == expected
        markdown = bridge.stage_export(plan, choose_path=True)
        assert markdown['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        for group in expected:
            for value in group:
                assert value + '  \n' in markdown['md']
        excel = bridge.export_excel('ladder', plan)
        assert excel['ok']
        book = load_workbook(excel['path'])
        try:
            for row, values in enumerate(expected, 2):
                assert book['配置阶梯'].cell(row, 5).value == '\n'.join(values)
        finally:
            book.close()
    finally:
        bridge.shutdown()
