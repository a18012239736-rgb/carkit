from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import stage_one
from engine.models import Cell, RawRow, RawTable, Trim


PLAN = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]


def raw_table(rows):
    count = len(next(iter(rows.values())))
    return RawTable(model='合成减少配置测试',
                    trims=[Trim(i, f'版型{i}', price_guide=10 + i) for i in range(count)],
                    rows=[RawRow(name=name, cells=cells) for name, cells in rows.items()])


def standard(values):
    return Cell(dot='●', text=values[0], subs=[Cell(dot='●', text=v) for v in values[1:]]) if values else Cell()


def changes(raw, target=1, base=0):
    return stage_one._compact_items(stage_one._change_items(stage_one.features(raw), target, base), base)


@pytest.mark.parametrize('name,before,after,expected', [
    ('主动闭合式进气格栅', Cell(dot='●'), Cell(), ['无（主动闭合式进气格栅）']),
    ('中控屏幕尺寸', Cell(text='15.6英寸'), Cell(text='10.25英寸'),
     ['10.25寸中控(15.6寸中控)', '10.25寸中控（15.6寸中控）']),
    ('中控屏幕尺寸', Cell(text='15.6英寸'), Cell(), ['无（15.6寸中控）']),
])
def test_entire_removal_and_scalar_downgrade_show_current_then_baseline(name, before, after, expected):
    assert changes(raw_table({name: [before, after]})) in [[value] for value in expected]


@pytest.mark.parametrize('before,after,expected', [
    (['CarPlay', 'HiCar'], ['CarPlay'], ['CarPlay（CarPlay+HiCar）']),
    (['CarPlay', 'HiCar'], ['HiCar', 'CarLife'], ['HiCar+CarLife（CarPlay+HiCar）']),
    (['CarPlay', 'HiCar'], [], ['无（CarPlay+HiCar）']),
    (['CarPlay', 'HiCar'], ['CarPlay', 'HiCar'], []),
    (['CarPlay'], ['CarPlay', 'HiCar'], ['HiCar']),
    ([], ['CarPlay'], ['CarPlay']),
])
def test_set_removals_are_single_lines_and_pure_additions_keep_existing_format(before, after, expected):
    raw = raw_table({'手机互联/映射': [standard(before), standard(after)]})
    assert changes(raw) == expected


def test_other_set_features_show_remaining_functions_not_only_removed_functions():
    raw = raw_table({'外后视镜功能': [standard(['电动调节', '电动折叠', '加热']),
                                   standard(['电动调节', '加热'])]})
    assert changes(raw) == ['外后视镜电调+外后视镜加热（外后视镜电调+外后视镜电动折叠+外后视镜加热）']


def test_comparison_uses_selected_baseline_not_neighboring_trim():
    raw = raw_table({'主动闭合式进气格栅': [Cell(dot='●'), Cell(), Cell()]})
    assert changes(raw, 2, 0) == ['无（主动闭合式进气格栅）']
    assert changes(raw, 2, 1) == []
    bridge = Bridge.__new__(Bridge)
    bridge.stage_raw = raw
    preview = bridge.stage_preview(PLAN + [{'target': 2, 'base': 0}])
    assert preview['ok']
    assert preview['columns'][2]['base'] == '版型0'
    assert preview['columns'][2]['items'] == ['无（主动闭合式进气格栅）']


@pytest.mark.parametrize('before,after,expected', [
    (['通风', '加热', '按摩'], ['加热'], '前排加热（前排通风加热按摩）'),
    (['通风', '加热'], [], '无（前排通风加热）'),
    (['通风'], ['加热'], '前排加热（前排通风）'),
])
def test_seat_compression_keeps_old_functions_inside_baseline_parentheses(before, after, expected):
    raw = raw_table({'前排座椅功能': [standard(before), standard(after)]})
    assert changes(raw) == [expected]


def test_standard_to_optional_still_lists_option_separately():
    raw = raw_table({'主动闭合式进气格栅': [Cell(dot='●'), Cell(dot='○')]})
    assert changes(raw) == ['无（主动闭合式进气格栅）']
    assert stage_one._option_lines(raw, 1) == ['主动闭合式进气格栅']
    config, optional = stage_one.render(raw, PLAN).split('## 选装')
    assert '无（主动闭合式进气格栅）' in config
    assert '> 选装：主动闭合式进气格栅' in optional


def test_preview_markdown_and_excel_share_removal_labels(tmp_path):
    raw = raw_table({
        '主动闭合式进气格栅': [Cell(dot='●'), Cell(dot='○')],
        '中控屏幕尺寸': [Cell(text='15.6英寸'), Cell()],
        '手机互联/映射': [standard(['CarPlay', 'HiCar']), standard(['HiCar', 'CarLife'])],
        '前排座椅功能': [standard(['通风', '加热', '按摩']), standard(['加热'])],
    })
    expected = ['无（主动闭合式进气格栅）', '无（15.6寸中控）',
                'HiCar+CarLife（CarPlay+HiCar）', '前排加热（前排通风加热按摩）']
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        preview = bridge.stage_preview(PLAN)
        assert preview['ok']
        assert preview['columns'][1]['items'] == expected
        assert preview['columns'][1]['options'] == ['主动闭合式进气格栅']
        markdown = bridge.stage_export(PLAN, choose_path=True)
        assert markdown['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        current = markdown['md'].split('**较版型0')[1].split('## 选装')[0]
        for label in expected:
            assert label + '  \n' in current
        assert '减少配置' not in current
        excel = bridge.export_excel('ladder', PLAN)
        assert excel['ok']
        book = load_workbook(excel['path'])
        try:
            sheet = book['配置阶梯']
            for row, column in enumerate(preview['columns'], 2):
                assert (sheet.cell(row, 5).value or '') == '\n'.join(column['items'])
                assert (sheet.cell(row, 6).value or '') == '\n'.join(column['options'])
        finally:
            book.close()
    finally:
        bridge.shutdown()
