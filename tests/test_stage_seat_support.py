"""Independent front-seat supports share preview/export data, never direction totals."""
from pathlib import Path

from openpyxl import load_workbook

from api.bridge import Bridge
from engine import seat_adjust, stage_one
from engine.mapper import map_raw_to_ladder
from engine.models import Cell, RawRow, RawTable, Trim
from engine.rules import Rules


PLAN = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]


def table(rows):
    return RawTable(model='座椅支持展示测试', trims=[Trim(0, '基础', price_guide=10),
                                             Trim(1, '升级', price_guide=11)],
                    rows=[RawRow(name=name, cells=cells) for name, cells in rows.items()])


def groups(raw, target=1, base=0, excluded=()):
    features = stage_one.features(raw, excluded)
    changes = stage_one._change_items(features, target, base)
    present, removed = stage_one._display_change_groups(features, changes, target, base)
    return stage_one._compact_items(present, base), stage_one._compact_items(removed, base)


def test_main_and_passenger_supports_have_independent_keys_without_extra_directions():
    raw = table({
        '主/副驾驶座电动调节': [Cell(text='主●/副●')] * 2,
        '主座椅调节方式': [Cell(dot='●', text='前后调节+靠背调节+高低调节(2向)',
            subs=[Cell(dot='●', text='腿托调节')])] * 2,
        '副座椅调节方式': [Cell(dot='●', text='前后调节+靠背调节',
            subs=[Cell(dot='●', text='腰部支撑(4向)')])] * 2,
    })
    original = raw.to_json()
    features = {feature['key']: feature for feature in stage_one.features(raw, [])}
    assert features['座椅腿托']['values'] == [['主驾腿托'], ['主驾腿托']]
    assert features['座椅腰托']['values'] == [['副驾腰托'], ['副驾腰托']]
    assert not features['座椅腿托']['scalar']
    assert not features['座椅腰托']['scalar']
    assert features['主驾座椅调节']['values'] == [['主驾6向电调'], ['主驾6向电调']]
    assert features['副驾座椅调节']['values'] == [['副驾4向电调'], ['副驾4向电调']]
    mapped = map_raw_to_ladder(raw, Rules())
    assert mapped[35]['values'] == ['主驾6向电调+副驾4向电调'] * 2
    assert seat_adjust.components(mapped[35]['values'][0]) == (10, 2)
    assert raw.to_json() == original


def test_optional_support_only_appears_in_optional_lines():
    raw = table({
        '主座椅调节方式': [Cell(dot='●', text='前后调节+靠背调节'),
            Cell(dot='●', text='前后调节+靠背调节',
                 subs=[Cell(dot='○', text='腿托调节'), Cell(dot='○', text='腰部支撑(4向)')])],
    })
    assert groups(raw) == ([], [])
    assert stage_one._option_lines(raw, 1, []) == ['主座椅调节方式：腿托调节 / 腰部支撑(4向)']
    assert not any(feature['key'] in {'座椅腿托', '座椅腰托'}
                   for feature in stage_one.features(raw, []))


def test_new_legrest_is_added_and_lost_lumbar_is_removed_below_the_baseline():
    raw = table({
        '主座椅调节方式': [Cell(text='前后调节+靠背调节'),
            Cell(text='前后调节+靠背调节', subs=[Cell(dot='●', text='腿托调节')])],
        '副座椅调节方式': [Cell(text='前后调节+靠背调节',
            subs=[Cell(dot='●', text='腰部支撑(4向)')]), Cell(text='前后调节+靠背调节')],
    })
    assert groups(raw) == (['主驾腿托'], ['副驾腰托'])


def test_one_to_both_seats_only_compares_the_new_or_lost_seat():
    raw = table({'前排座椅腰托': [Cell(dot='●', text='主驾'), Cell(dot='●', text='前排')]})
    features = {feature['key']: feature['values'] for feature in stage_one.features(raw, [])}
    assert features['座椅腰托'] == [['主驾腰托'], ['主驾腰托', '副驾腰托']]
    assert groups(raw, target=0, base=None) == (['主驾腰托'], [])
    assert groups(raw, target=1, base=None) == (['主驾腰托+副驾腰托'], [])
    assert groups(raw) == (['副驾腰托'], [])
    assert groups(raw, target=0, base=1) == ([], ['副驾腰托'])


def test_supports_in_front_seat_function_rows_are_not_repeated_as_unknown_features():
    raw = table({'前排座椅功能': [Cell(text='加热(仅驾驶位)'),
        Cell(dot='●', text='加热(仅驾驶位)', subs=[
            Cell(dot='●', text='腿托(仅主驾)'), Cell(dot='●', text='腰部支撑(仅副驾)')])]})
    features = {feature['key']: feature['values'] for feature in stage_one.features(raw, [])}
    assert features['前排座椅功能'] == [['主驾座椅加热'], ['主驾座椅加热']]
    assert features['座椅腿托'] == [[], ['主驾腿托']]
    assert features['座椅腰托'] == [[], ['副驾腰托']]
    assert groups(raw) == (['主驾腿托', '副驾腰托'], [])


def test_original_source_row_exclusions_control_supports_and_optional_lines():
    raw = table({
        '主座椅调节方式': [Cell(text='前后调节+靠背调节'),
            Cell(text='前后调节+靠背调节', subs=[Cell(dot='●', text='腿托调节'),
                Cell(dot='○', text='腰部支撑(4向)')])],
        '副座椅调节方式': [Cell(text='前后调节+靠背调节'),
            Cell(text='前后调节+靠背调节', subs=[Cell(dot='●', text='腰部支撑(4向)')])],
    })
    assert groups(raw) == (['主驾腿托', '副驾腰托'], [])
    assert groups(raw, excluded=['主座椅调节方式']) == (['副驾腰托'], [])
    assert groups(raw, excluded=['副座椅调节方式']) == (['主驾腿托'], [])
    assert groups(raw, excluded=['主座椅调节方式', '副座椅调节方式']) == ([], [])
    assert stage_one._option_lines(raw, 1, ['主座椅调节方式']) == []
    front = table({'前排座椅功能': [Cell(), Cell(text='腿托')]})
    assert groups(front) == (['主驾腿托+副驾腿托'], [])
    assert groups(front, excluded=['前排座椅功能']) == ([], [])


def test_preview_markdown_and_excel_keep_the_same_independent_support_items(tmp_path):
    raw = table({
        '主/副驾驶座电动调节': [Cell(text='主●/副●')] * 2,
        '主座椅调节方式': [Cell(text='前后调节+靠背调节+高低调节(2向)'),
            Cell(text='前后调节+靠背调节+高低调节(2向)', subs=[Cell(dot='●', text='腿托调节')])],
        '副座椅调节方式': [Cell(text='前后调节+靠背调节',
            subs=[Cell(dot='●', text='腰部支撑(4向)')]), Cell(text='前后调节+靠背调节')],
    })
    expected = [['主驾6向电调+副驾4向电调', '副驾腰托'], ['主驾腿托', '无（副驾腰托）']]
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        assert bridge.save_stage_display_rules([])['ok']
        preview = bridge.stage_preview(PLAN)
        assert preview['ok']
        assert [column['items'] for column in preview['columns']] == expected
        assert preview['columns'][1]['display_items'] == ['主驾腿托']
        assert preview['columns'][1]['removed_items'] == ['副驾腰托']
        markdown = bridge.stage_export(PLAN, choose_path=True)
        assert markdown['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        for items in expected:
            for value in items:
                assert value + '  \n' in markdown['md']
        excel = bridge.export_excel('ladder', PLAN)
        assert excel['ok']
        book = load_workbook(excel['path'])
        try:
            assert book['配置阶梯'].cell(2, 5).value == '主驾6向电调+副驾4向电调\n副驾腰托'
            assert book['配置阶梯'].cell(3, 5).value == '主驾腿托\n无（副驾腰托）'
        finally:
            book.close()
    finally:
        bridge.shutdown()
