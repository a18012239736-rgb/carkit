import json
from pathlib import Path

import pytest

from api.bridge import Bridge
from engine import storage
from engine.models import (Cell, DiffResult, Ladder, LadderItem, RawRow,
                           RawTable, Snapshot, Trim, ValuationTable)
from engine.snapshot import prepare_snapshot
from engine.valuer import default_valuation


@pytest.mark.parametrize('model', [RawTable(), Ladder(), Snapshot(), DiffResult(), ValuationTable()])
def test_model_save_replaces_complete_json_without_temporary_files(tmp_path, model):
    path = tmp_path / 'data.json'
    storage.atomic_write_text(path, '{"old":true}')
    model.save(path)
    assert json.loads(path.read_text(encoding='utf-8')) == json.loads(model.to_json())
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize('failure', ['fsync', 'replace'])
def test_failed_save_preserves_existing_file_and_cleans_temporary_file(tmp_path, monkeypatch, failure):
    path = tmp_path / 'data.json'
    storage.atomic_write_text(path, '{"old":true}')

    def fail(*args):
        raise OSError('模拟保存中断')

    monkeypatch.setattr(storage.os, failure, fail)
    with pytest.raises(OSError, match='保存中断'):
        Snapshot(model='新数据').save(path)
    assert json.loads(path.read_text(encoding='utf-8')) == {'old': True}
    assert list(tmp_path.iterdir()) == [path]


def _raw(year):
    return RawTable(model='测试车型', series_id='7538',
                    trims=[Trim(idx=0, short='基础版', full=f'测试车型 {year}款 基础版', price_guide=10)],
                    rows=[RawRow(name='厂商指导价(元)', cells=[Cell(text='10万')]),
                          RawRow(name='CLTC纯电续航里程(km)', cells=[Cell(text='500')])])


def test_repeated_capture_and_import_keep_both_histories(tmp_path):
    bridge = Bridge(str(tmp_path / 'work'))
    try:
        first = bridge._stage_store(_raw(2025).to_dict())
        second = bridge._stage_store(_raw(2027).to_dict())
        assert first['raw_path'] != second['raw_path']
        assert '2025款' in RawTable.load(first['raw_path']).trims[0].full
        assert '2027款' in RawTable.load(second['raw_path']).trims[0].full
        history = bridge.vehicle_history()
        assert len(history) == 2
        assert all('1个版型' in record['label'] and '日期未记录' not in record['label'] for record in history)

        source = tmp_path / 'source.json'
        _raw(2027).save(source)
        a, b = bridge.import_raw(str(source)), bridge.import_raw(str(source))
        assert a['ok'] and b['ok'] and a['path'] != b['path']
        assert len(bridge.vehicle_history()) == 4
    finally:
        bridge.shutdown()


def test_failed_valuation_save_preserves_last_saved_rules(tmp_path, monkeypatch):
    bridge = Bridge(str(tmp_path))
    try:
        initial = bridge.save_current_valuation(default_valuation(bridge.rules))
        path = Path(initial['path'])
        original = path.read_bytes()
        data = initial['valuation']
        next(item for item in data['items'] if item['no'] == 40)['val'] = 9999

        def fail(*args):
            raise OSError('模拟保存中断')

        monkeypatch.setattr(storage.os, 'replace', fail)
        assert not bridge.save_current_valuation(data)['ok']
        assert path.read_bytes() == original
        assert len(list(path.parent.iterdir())) == 1
    finally:
        bridge.shutdown()


def test_bridge_uses_saved_rule_direction_and_keeps_repeated_results(tmp_path):
    bridge = Bridge(str(tmp_path))
    try:
        left, right = [], []
        for item in bridge.rules.items:
            no = item['no']
            left.append(LadderItem(no=no, name=item['name'], values=['R17钢轮毂' if no == 4 else '✕']))
            right.append(LadderItem(no=no, name=item['name'], values=['R16铝轮毂' if no == 4 else '✕']))
        left_path, right_path = tmp_path / 'left.json', tmp_path / 'right.json'
        Ladder(model='本品', trims=[{'name': 'A', 'price_guide': 10}], items=left).save(left_path)
        Ladder(model='竞品', trims=[{'name': 'B', 'price_guide': 10}], items=right).save(right_path)
        rules = default_valuation(bridge.rules)
        next(item for item in rules['items'] if item['no'] == 4)['params'].update(per_inch=100, alloy_bonus=500)
        assert bridge.save_current_valuation(rules)['ok']
        pairs = [{'self_trim': 'A', 'comp_trim': 'B'}]
        a = bridge.run_diff(str(left_path), str(right_path), pairs, left_is_competitor=True)
        b = bridge.run_diff(str(left_path), str(right_path), pairs, left_is_competitor=True)
        assert a['ok'] and b['ok'], (a, b)
        assert a['groups'][0]['valuation']['config_adv'] == -400
        assert a['md_path'] != b['md_path']
        assert Path(a['md_path']).read_text(encoding='utf-8') == a['md']
        assert Path(b['md_path']).read_text(encoding='utf-8') == b['md']
    finally:
        bridge.shutdown()


def test_incomplete_snapshot_can_add_and_normalize_default_usb_row():
    snapshot = prepare_snapshot(Snapshot(trims=[{'name': '基础版'}]))
    assert next(cell for cell in snapshot.cells if cell['no'] == 32)['values']['基础版'] == 'USB/Type-C 3个'


def test_cli_defaults_keep_history_and_do_not_overwrite_adjacent_rule_json(tmp_path, monkeypatch):
    from argparse import Namespace
    from openpyxl import load_workbook
    from engine.rules import Rules
    from engine.valuer import make_template_xlsx
    from engine.excel_export import export_report
    from cli import cmd_diff, cmd_ladder, cmd_snapshot

    monkeypatch.chdir(tmp_path)
    raw_path = tmp_path / 'raw.json'
    _raw(2027).save(raw_path)
    args = Namespace(raw=str(raw_path), o=None, md=None, series='', model='')
    cmd_ladder(args)
    cmd_ladder(args)
    assert len(list(tmp_path.glob('竞品阶梯-*.json'))) == 2

    rules = Rules()
    snapshot = prepare_snapshot(Snapshot(model='本品', trims=[{'name': 'A', 'price_guide': 10}],
        cells=[{'no': item['no'], 'values': {'A': 'R17钢轮毂' if item['no'] == 4 else '●' if item['no'] == 40 else '✕'}}
               for item in rules.items]))
    snapshot_path = tmp_path / 'source.json'
    snapshot.save(snapshot_path)
    cmd_snapshot(Namespace(json=str(snapshot_path), o=None))
    cmd_snapshot(Namespace(json=str(snapshot_path), o=None))
    assert len(list(tmp_path.glob('本品-配置快照-*.json'))) == 2

    ladder = snapshot.to_ladder(rules.items)
    ladder.model = '竞品'
    ladder.item(4).values = ['R16铝轮毂']
    ladder.item(40).values = ['✕']
    ladder_path = tmp_path / 'right.json'
    ladder.save(ladder_path)
    data = default_valuation(rules)
    by_no = {item['no']: item for item in data['items']}
    by_no[4]['params'].update(per_inch=100, alloy_bonus=500)
    by_no[40]['rule'] = 'excluded'
    xlsx = tmp_path / 'rules.xlsx'
    make_template_xlsx(xlsx, rules, by_no)
    adjacent = tmp_path / 'rules.json'
    storage.atomic_write_text(adjacent, '{"existing":true}')
    result_path = tmp_path / 'result.json'
    args = Namespace(snapshot=str(snapshot_path), ladder=str(ladder_path),
                     pairs='A:A', valuation=str(xlsx), o=None, json=str(result_path))
    cmd_diff(args)
    cmd_diff(args)
    assert len(list(tmp_path.glob('赋值对比-*.md'))) == 2
    assert json.loads(adjacent.read_text(encoding='utf-8')) == {'existing': True}
    result = json.loads(result_path.read_text(encoding='utf-8'))
    assert 40 not in {cell['no'] for cell in result['cells']}
    assert result['groups'][0]['valuation']['config_adv'] == -400

    report = tmp_path / 'report.xlsx'
    export_report(report, 'diff', result)
    book = load_workbook(report)
    assert book['对比汇总']['B5'].value == -400
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    assert cell.number_format == '0.##;[Red]-0.##'
    assert next(row[5] for row in book['赋值明细'].iter_rows(min_row=2, values_only=True) if 'R17' in row[3]) == -400
    book.close()
