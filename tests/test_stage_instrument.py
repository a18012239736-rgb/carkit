from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import stage_one
from engine.mapper import map_raw_to_ladder
from engine.models import Cell, RawRow, RawTable, Trim
from engine.rules import Rules
from engine.valuer import dynamic_delta


def table(full, sizes, reverse=False):
    rows = [RawRow('全液晶仪表盘', full), RawRow('液晶仪表尺寸', sizes)]
    return RawTable(model='仪表展示测试', trims=[Trim(i, f'版型{i}', price_guide=10+i) for i in range(len(full))],
                    rows=rows[::-1] if reverse else rows)


@pytest.mark.parametrize('reverse', [False, True])
def test_full_lcd_and_size_merge_in_either_source_order_without_affecting_calculation(reverse):
    raw = table([Cell(dot='●'), Cell(dot='●')], [Cell(dot='●', text='8.8英寸'), Cell(text='10.17英寸')], reverse)
    original = raw.to_json()
    mapped = map_raw_to_ladder(raw, Rules())
    amount = dynamic_delta(29, mapped[29]['values'][1], mapped[29]['values'][0])
    assert stage_one.features(raw, []) == [{'key': '仪表', 'values': [['全液晶8.8寸仪表'], ['全液晶10.17寸仪表']], 'scalar': True}]
    assert raw.to_json() == original
    after = map_raw_to_ladder(raw, Rules())
    assert after == mapped
    assert dynamic_delta(29, after[29]['values'][1], after[29]['values'][0]) == amount


def test_instrument_size_or_full_lcd_changes_keep_current_and_baseline_descriptions():
    raw = table([Cell(dot='●'), Cell(dot='●'), Cell(), Cell(dot='●')],
                [Cell(text='8.8英寸'), Cell(text='10.17英寸'), Cell(text='8.8英寸'), Cell()])
    features = stage_one.features(raw, [])
    assert stage_one._change_items(features, 1, 0) == [('仪表', '全液晶10.17寸仪表(全液晶8.8寸仪表)')]
    assert stage_one._change_items(features, 2, 0) == [('仪表', '8.8寸仪表(全液晶8.8寸仪表)')]
    assert stage_one._change_items(features, 3, 0) == [('仪表', '全液晶仪表(全液晶8.8寸仪表)')]


def test_optional_lcd_or_size_is_not_added_to_standard_instrument():
    raw = table([Cell(dot='○'), Cell(dot='●'), Cell(dot='○')],
                [Cell(dot='●', text='10.17英寸'), Cell(dot='○', text='12.3英寸'), Cell(dot='○', text='8.8英寸')])
    assert stage_one.features(raw, [])[0]['values'] == [['10.17寸仪表'], ['全液晶仪表'], []]
    assert stage_one._option_lines(raw, 0, []) == ['全液晶仪表盘']
    assert stage_one._option_lines(raw, 1, []) == ['液晶仪表尺寸：12.3寸']
    assert stage_one._option_lines(raw, 2, []) == ['全液晶仪表盘', '液晶仪表尺寸：8.8寸']


@pytest.mark.parametrize('excluded, expected', [
    ([], ['全液晶8.8寸仪表']),
    (['全液晶仪表盘'], ['8.8寸仪表']),
    (['液晶仪表尺寸'], ['全液晶仪表']),
    (['全液晶仪表盘', '液晶仪表尺寸'], None),
])
def test_each_instrument_source_can_be_hidden_independently(excluded, expected):
    raw = table([Cell(dot='●')], [Cell(dot='●', text='8.8英寸')])
    features = stage_one.features(raw, excluded)
    assert features == ([] if expected is None else [{'key': '仪表', 'values': [expected], 'scalar': True}])


def test_preview_markdown_and_excel_share_instrument_summary_and_separate_options(tmp_path):
    raw = table([Cell(dot='●'), Cell(dot='●'), Cell(dot='○')],
                [Cell(text='8.8英寸'), Cell(text='10.17英寸'), Cell(text='10.17英寸')])
    plan = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}, {'target': 2, 'base': 1}]
    expected = ['全液晶8.8寸仪表', '全液晶10.17寸仪表(全液晶8.8寸仪表)', '10.17寸仪表(全液晶10.17寸仪表)']
    original = raw.to_json()
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        assert bridge.save_stage_display_rules([])['ok']
        preview = bridge.stage_preview(plan)
        assert preview['ok']
        assert [column['items'] for column in preview['columns']] == [[value] for value in expected]
        assert [column['options'] for column in preview['columns']] == [[], [], ['全液晶仪表盘']]
        markdown = bridge.stage_export(plan, choose_path=True)
        assert markdown['ok']
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        for value in expected:
            assert value + '  \n' in markdown['md']
        assert '> 选装：全液晶仪表盘' in markdown['md']
        excel = bridge.export_excel('ladder', plan)
        assert excel['ok']
        book = load_workbook(excel['path'])
        try:
            for row, column in enumerate(preview['columns'], 2):
                assert book['配置阶梯'].cell(row, 5).value == '\n'.join(column['items'])
                assert (book['配置阶梯'].cell(row, 6).value or '') == '\n'.join(column['options'])
        finally:
            book.close()
        assert raw.to_json() == original
    finally:
        bridge.shutdown()
