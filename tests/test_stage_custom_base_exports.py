from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine.models import Cell, RawRow, RawTable, Trim


PLAN = [
    {'target': 1, 'base': None},
    {'target': 0, 'base': 1},
    {'target': 2, 'base': 0},
]


def raw_table(trim_ids=(0, 1, 2)):
    return RawTable(model='自选基本版型测试车', source='rich', trims=[
        Trim(trim_ids[0], '低价型', price_guide=10),
        Trim(trim_ids[1], '高价型', price_guide=20),
        Trim(trim_ids[2], '中价型', price_guide=15),
    ], rows=[
        RawRow('共同功能', [Cell(dot='●') for _ in trim_ids]),
        RawRow('智能香氛', [Cell(), Cell(dot='●'), Cell(dot='●')]),
        RawRow('电动遮阳帘', [Cell(dot='●'), Cell(), Cell()]),
        RawRow('舒适套装', [Cell(dot='○', text=name) for name in ['低价包', '高价包', '中价包']]),
    ])


@pytest.mark.parametrize('trim_ids', [(0, 1, 2), (101, 303, 202)])
def test_high_price_nonzero_base_has_consistent_preview_and_exports(tmp_path, trim_ids):
    raw = raw_table(trim_ids)
    original = raw.to_json()
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path / name)
    try:
        preview = bridge.stage_preview(PLAN)
        markdown = bridge.stage_export(PLAN)
        excel = bridge.export_excel('ladder', PLAN)
        assert preview['ok'] and markdown['ok'] and excel['ok']
        columns = preview['columns']
        assert [column['name'] for column in columns] == ['高价型', '低价型', '中价型']
        assert [column['base'] for column in columns] == [None, '高价型', '低价型']
        assert [column['price'] for column in columns] == [20, 10, 15]
        assert columns[0]['items'] == ['共同功能', '智能香氛']
        assert columns[1]['items'] == ['无（智能香氛）', '电动遮阳帘']
        assert columns[2]['items'] == ['智能香氛', '无（电动遮阳帘）']
        assert [column['options'] for column in columns] == [
            ['舒适套装：高价包'], ['舒适套装：低价包'], ['舒适套装：中价包'],
        ]

        md = markdown['md']
        assert Path(markdown['path']).read_text(encoding='utf-8') == md
        table_lines = [line for line in md.splitlines() if line.startswith('| ')][1:]
        assert table_lines == [
            '| 高价型 | 20 | 基本配置 |',
            '| 低价型 | 10 | 高价型 |',
            '| 中价型 | 15 | 低价型 |',
        ]
        assert md.index('**基本配置：高价型') < md.index('**较高价型 -10.00万:(低价型)')
        assert md.index('**较高价型 -10.00万:(低价型)') < md.index('**较低价型 +5.00万:(中价型)')
        for column in columns:
            assert all(item in md for item in column['items'])
            assert all('> 选装：' + option in md for option in column['options'])

        book = load_workbook(excel['path'])
        try:
            rows = list(book['配置阶梯'].iter_rows(min_row=2, values_only=True))
            assert rows == [
                (raw.model, column['name'], column['price'], column['base'] or '基础配置',
                 '\n'.join(column['items']), '\n'.join(column['options']))
                for column in columns
            ]
        finally:
            book.close()
        assert bridge.stage_raw is raw
        assert raw.to_json() == original
    finally:
        bridge.shutdown()


@pytest.mark.parametrize('plan', [
    [{'target': 1, 'base': 1}],
    [{'target': 1, 'base': None}, {'target': 0, 'base': 2}],
    [{'target': 1, 'base': 0}, {'target': 0, 'base': 1}],
    [{'target': 1, 'base': None}, {'target': 1, 'base': 1}],
])
def test_invalid_custom_base_is_rejected_by_all_export_paths(tmp_path, plan):
    bridge = Bridge(str(tmp_path / 'work'))
    bridge.stage_raw = raw_table()
    original = bridge.stage_raw.to_json()
    dialogs = []
    bridge.save_file_dialog = lambda *args: dialogs.append(args)
    try:
        results = [bridge.stage_preview(plan), bridge.stage_export(plan), bridge.export_excel('ladder', plan)]
        assert all(result['ok'] is False and result.get('error') for result in results)
        assert not dialogs
        assert not list(tmp_path.rglob('*.md'))
        assert not list(tmp_path.rglob('*.xlsx'))
        assert bridge.stage_raw.to_json() == original
    finally:
        bridge.shutdown()
