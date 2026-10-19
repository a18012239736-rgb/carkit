import hashlib
from pathlib import Path

from openpyxl import load_workbook
import pytest

from api.bridge import Bridge
from engine import stage_one
from engine.differ import diff
from engine.ladder import build_ladder
from engine.models import Cell, Ladder, RawRow, RawTable, Trim, ValuationItem, ValuationTable
from engine.rules import Rules
from engine.valuer import default_valuation, dynamic_delta, value_pair


AR_ROW = 'AR-HUD增强现实抬头显示'
PLAN = [{'target': 0, 'base': None}, {'target': 1, 'base': 0}]


def _raw(rows, count=2):
    return RawTable(model='HUD合成测试车', trims=[
        Trim(idx=i, short=f'版型{i}', price_guide=10+i) for i in range(count)
    ], rows=[RawRow(name=name, cells=cells) for name, cells in rows.items()])


def test_screenshot_ar_hud_flows_through_mapping_comparison_preview_and_exports(tmp_path):
    # Presence and size reproduce the user's A06 screenshot; no size-based pricing.
    raw = _raw({AR_ROW: [Cell(text='-'), Cell(dot='●')],
                'HUD抬头尺寸': [Cell(text='-'), Cell(dot='●', text='49英寸')]})
    rules = Rules()
    ladder = build_ladder(raw, rules)
    assert ladder.item(30).values == ['✕', '●AR-HUD']
    assert not ladder.item(30).row_absent
    cells = diff(ladder, ladder, [{'self_trim': '版型1', 'comp_trim': '版型0'}], rules)
    hud = next(cell for cell in cells if cell['no'] == 30)
    assert hud['verdict'] == '多'
    assert 'AR-HUD' in hud['backup_more']
    valuation = ValuationTable(items=[ValuationItem(**item) for item in default_valuation(rules)['items']])
    assert value_pair({}, {}, 10, 10, valuation, cells=cells)['config_adv'] == 2000

    bridge = Bridge(str(tmp_path/'work'))
    bridge.stage_raw = raw
    bridge.save_file_dialog = lambda name, *_: str(tmp_path/name)
    try:
        assert bridge.save_stage_display_rules([])['ok']
        preview = bridge.stage_preview(PLAN)
        assert preview['ok'], preview
        # Clearing extra exclusions cannot reveal the code-excluded size row.
        assert preview['columns'][1]['items'] == [AR_ROW]
        markdown = bridge.stage_export(PLAN, choose_path=True)
        assert markdown['ok'], markdown
        assert Path(markdown['path']).read_text(encoding='utf-8') == markdown['md']
        assert AR_ROW in markdown['md']
        assert 'HUD抬头尺寸' not in markdown['md']
        excel = bridge.export_excel('ladder', PLAN)
        assert excel['ok'], excel
        book = load_workbook(excel['path'])
        try:
            assert book['配置阶梯'].cell(3, 5).value == AR_ROW
        finally:
            book.close()
    finally:
        bridge.shutdown()


@pytest.mark.parametrize('name,cell,expected', [
    (AR_ROW, Cell(dot='●'), '●AR-HUD'),
    ('AR-HUD', Cell(dot='●'), '●AR-HUD'),
    ('HUD抬头数字显示', Cell(dot='●', text='P-HUD'), '●P-HUD'),
    ('HUD', Cell(dot='●'), '●HUD'),
    ('HUD', Cell(subs=[Cell(dot='●', text='AR-HUD'), Cell(dot='○', text='P-HUD')]), '●AR-HUD'),
    (AR_ROW, Cell(dot='●', text='[待定]'), '[待定]'),
])
def test_hud_type_uses_row_name_and_only_standard_subitems(name, cell, expected):
    raw = _raw({name: [cell]}, count=1)
    assert build_ladder(raw, Rules()).item(30).values == [expected]


@pytest.mark.parametrize('reverse_rows', [False, True])
def test_separate_ar_row_is_not_masked_by_generic_hud_or_promoted_from_optional(reverse_rows):
    rows = {'HUD抬头数字显示': [Cell(), Cell(dot='●')],
            AR_ROW: [Cell(dot='●'), Cell(dot='○', text='(3000元)')]}
    if reverse_rows:
        rows = dict(reversed(list(rows.items())))
    raw = _raw(rows)
    assert build_ladder(raw, Rules()).item(30).values == ['●AR-HUD', '●HUD']
    assert stage_one._option_lines(raw, 1) == [AR_ROW+'：(3000元)']


def test_optional_ar_and_size_never_become_standard_equipment():
    raw = _raw({AR_ROW: [Cell(dot='○', text='(3000元)'), Cell()],
                'HUD抬头尺寸': [Cell(dot='○', text='49英寸'), Cell(dot='●', text='49英寸')]})
    values = build_ladder(raw, Rules()).item(30).values
    assert values == ['○AR-HUD(选装)', '✕']
    assert dynamic_delta(30, values[0], '✕') == 0
    features = stage_one.features(raw, excluded_rows=[])
    assert stage_one._change_items(features, 0, None) == []
    # The full ladder shows the explicitly standard size as a separate source
    # row, but it must never infer a standard HUD or price it as one.
    assert stage_one._compact_items(stage_one._change_items(features, 1, None), None) == ['HUD抬头尺寸49寸']
    assert AR_ROW+'：(3000元)' in stage_one._option_lines(raw, 0)


@pytest.mark.parametrize('generic_dot', [None, '', '●'])
def test_old_hud_cache_updates_only_untouched_values_by_trim_name(tmp_path, generic_dot):
    rows = {AR_ROW: [Cell(dot='●'), Cell(dot='●'), Cell(dot='○', text='(3000元)')]}
    if generic_dot is not None:
        rows['HUD抬头数字显示'] = [Cell(dot=generic_dot) for _ in range(3)]
    raw = _raw(rows, count=3)
    bridge = Bridge(str(tmp_path))
    try:
        capture = tmp_path/'raw'/'capture.json'
        raw.save(str(capture))
        capture_bytes = capture.read_bytes()
        key = hashlib.sha256(capture_bytes+b'|cell-marker-v2').hexdigest()[:16]
        path = tmp_path/'阶梯'/('compare-'+key+'.json')
        cached = build_ladder(raw, bridge.rules)
        # Freeze the old mapper's actual outputs, independently of the new code.
        old_value = '●HUD' if generic_dot == '●' else '✕'
        cached.item(30).values = [old_value, '●P-HUD（人工修正）', old_value]
        cached.item(30).row_absent = generic_dot is None
        cached.item(17).values = ['人工灯光0', '人工灯光1', '人工灯光2']
        cached.trims[1]['price_guide'] = 17.68
        cached.trims.reverse()
        for item in cached.items:
            item.values.reverse()
            for sub in item.subs:
                sub['values'].reverse()
        cached.save(str(path))

        result = bridge.prepare_competitor('capture.json')
        assert result['ok'], result
        assert result['path'] == str(path)
        migrated = Ladder.from_dict(result['ladder'])
        optional_or_standard = '●HUD' if generic_dot == '●' else '○AR-HUD(选装)'
        assert migrated.item(30).values == [optional_or_standard, '●P-HUD（人工修正）', '●AR-HUD']
        assert not migrated.item(30).row_absent
        assert migrated.trims == cached.trims
        assert migrated.item(17).values == cached.item(17).values
        assert capture.read_bytes() == capture_bytes
        migrated_bytes = path.read_bytes()
        again = bridge.prepare_competitor('capture.json')
        assert again['ok'] and again['ladder'] == result['ladder']
        assert path.read_bytes() == migrated_bytes
    finally:
        bridge.shutdown()
